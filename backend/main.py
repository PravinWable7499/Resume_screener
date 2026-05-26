import asyncio
import io
import json
import os
import random
import re
import secrets
import time
import xml.etree.ElementTree as ET
import zipfile
from datetime import date, datetime, timedelta
from email.message import EmailMessage
from typing import List, Optional

import aiosmtplib
import groq
import pdfplumber
import docx
from dotenv import load_dotenv
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel


from database import (
    Candidate as DBCandidate,
    CandidateFeedback,
    ClientLearning,
    DailyContent,
    DailyGame,
    FeatureUsage,
    GameResult,
    Invitation,
    MonthlySurvey,
    OTPCode,
    Project as DBProject,
    Recruiter,
    SessionLocal,
    Streak,
    create_tables,
)

load_dotenv()

app = FastAPI(title="ResumeAI API")

app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:3000",
        "http://localhost:5173",
        "https://resumeai.agile-tech.in",
        "https://www.resumeai.agile-tech.in",
    ],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

client = groq.Groq()  # reads GROQ_API_KEY from env

ADMIN_EMAIL = os.getenv("ADMIN_EMAIL", "chetan@agile-tech.in")


@app.get("/health")
def health_check():
    return {"status": "ok", "app": "ResumeAI", "version": "2.0"}


@app.on_event("startup")
async def _startup():
    create_tables()
    db = SessionLocal()
    try:
        # Merge duplicate recruiter records caused by case-insensitive email matches
        all_recruiters = db.query(Recruiter).all()
        groups: dict = {}
        for r in all_recruiters:
            key = r.email.lower().strip()
            groups.setdefault(key, []).append(r)

        for key, group in groups.items():
            if len(group) > 1:
                proj_counts = {
                    r.id: db.query(DBProject).filter(DBProject.created_by_email == r.email).count()
                    for r in group
                }
                group.sort(
                    key=lambda r: (int(r.is_admin or False), proj_counts[r.id]),
                    reverse=True,
                )
                keeper = group[0]
                keeper.email = key
                for dup in group[1:]:
                    db.query(DBProject).filter(DBProject.created_by_email == dup.email).update(
                        {"created_by_email": key}
                    )
                    db.query(OTPCode).filter(OTPCode.email == dup.email).update(
                        {"email": key}
                    )
                    db.delete(dup)
            elif group[0].email != key:
                group[0].email = key

        db.commit()

        # Ensure admin account has is_admin flag and role='admin'
        admin_email = ADMIN_EMAIL.lower().strip()
        admin = db.query(Recruiter).filter(Recruiter.email == admin_email).first()
        if admin:
            changed = False
            if not admin.is_admin:
                admin.is_admin = True
                changed = True
            if admin.role != 'admin':
                admin.role = 'admin'
                changed = True
            if changed:
                db.commit()

        # Clean up duplicate candidates (same filename within a project)
        try:
            _deduplicate_candidates(db)
        except Exception:
            pass
    finally:
        db.close()


def _deduplicate_candidates(db) -> int:
    projects = db.query(DBProject).all()
    total_deleted = 0
    for project in projects:
        candidates = (
            db.query(DBCandidate)
            .filter(DBCandidate.project_id == project.id)
            .all()
        )
        by_filename: dict = {}
        for c in candidates:
            by_filename.setdefault(c.filename, []).append(c)
        for group in by_filename.values():
            if len(group) <= 1:
                continue
            group.sort(key=lambda c: c.id, reverse=True)  # highest id = most recent
            for dup in group[1:]:
                db.delete(dup)
                total_deleted += 1
    if total_deleted:
        db.commit()
    return total_deleted


# ─── Request / Response Models ────────────────────────────────────────────────

class OTPRequest(BaseModel):
    email: str
    name: str = ""

class OTPVerify(BaseModel):
    email: str
    code: str

class ProjectCreate(BaseModel):
    name: str
    client_name: str = ""
    created_by_email: str

class JDContextUpdate(BaseModel):
    jd_text: str = ""
    jd_context: str = ""
    required_location: str = ""
    keywords: str = ""

class CandidateSave(BaseModel):
    name: str = ""
    filename: str = ""
    score: int = 0
    result_json: str = "{}"
    candidate_phone: str | None = None

class CallGuideRequest(BaseModel):
    jd_text: str
    candidate_json: dict

class WhatsAppMessageRequest(BaseModel):
    message_type: str = "check_interest"  # "check_interest" | "other"
    context: str = ""
    recruiter_name: str = "Recruiter"

class RemarkGenerate(BaseModel):
    linkedin_url: str
    linkedin_text: str
    voice_transcript: str = ""

class RemarkSave(BaseModel):
    remark: str
    linkedin_url: str = ""
    linkedin_text: str = ""

class CandidateStatusUpdate(BaseModel):
    status: str

class PipelineStageUpdate(BaseModel):
    stage: str

class ChatMessage(BaseModel):
    role: str
    content: str

class ChatRequest(BaseModel):
    candidate_id: int
    message: str
    chat_history: list[ChatMessage] = []

class InviteRequest(BaseModel):
    email: str
    name: str = ""
    role: str = "recruiter"
    team_manager_email: str = ""
    requester_email: str

class ActivateRequest(BaseModel):
    token: str

class RecruiterUpdate(BaseModel):
    status: str | None = None
    role: str | None = None

class GameResultRequest(BaseModel):
    email: str
    game_type: str
    score: int
    completed: bool
    time_taken: int = 0

class FeedbackRequest(BaseModel):
    project_id: int
    recruiter_email: str
    action: str
    reason: str
    reason_detail: str = ""

class FeatureUsageRequest(BaseModel):
    recruiter_email: str
    feature_name: str

class SurveyRequest(BaseModel):
    recruiter_email: str
    q1_rating: int = 0
    q2_rating: int = 0
    q3_text: str = ""

_semaphore = asyncio.Semaphore(4)

SYSTEM_PROMPT = (
    "You are an expert ATS (Applicant Tracking System) and senior HR recruiter. "
    "Evaluate resumes against job descriptions objectively and provide structured JSON insights. "
    "Respond with valid JSON only — no markdown, no code fences, no extra text."
)

MAX_RESUME_CHARS = 8000

_JSON_SCHEMA = """{
  "candidate_name": "Full name extracted from resume",
  "score": <integer 0-100>,
  "experience_total_months": <integer: sum of all job tenure_months from actual start/end dates — never from a summary line. CRITICAL: always floor to complete months — do NOT round up or round to nearest. Example: started June 17 2024, today May 20 2026 = 23 months (1 year 11 months and 3 days — the 3 partial days are ignored). Apply the same floor rule to every individual job tenure before summing>,
  "experience_relevant_months": <integer: months in roles whose responsibilities directly match the JD. CRITICAL: same floor rule — complete months only, never round up partial months>,
  "total_experience_years": <float: experience_total_months ÷ 12 rounded to 1 decimal — e.g. 3.3, 1.9, 0.7. CRITICAL: must be the precise float — 1 year 11 months = 1.9, NOT 1.0>,
  "total_experience_display": <string — CRITICAL: ALWAYS include months if there are remaining months after the years. Example: a candidate who started June 2024 and it is now May 2026 has 1 year and 11 months — display MUST be "1 yr 11 mo" NOT "1 yr". Never round down to whole years. Never drop months. Calculate precisely from actual start and end dates. Format rules: if remainder months = 0 show years only ("2 yrs"); if total < 12 months show months only ("8 mo"); otherwise ALWAYS show both ("1 yr 11 mo", "3 yrs 4 mo"). Use "yr" for exactly 1 year, "yrs" for 2+, always "mo" for months>,
  "relevant_experience_years": <float: experience_relevant_months ÷ 12 rounded to 1 decimal. CRITICAL: precise float — never truncate months>,
  "relevant_experience_display": <string: same CRITICAL rules as total_experience_display — always include months if non-zero remainder>,
  "current_role": "Most recent or current job title",
  "birth_year": <4-digit birth year if mentioned anywhere on resume, else null>,
  "education_level": <EXACTLY one of: "10th", "12th", "Diploma", "BSc/BA/BCom", "BE/BTech", "PG/MBA/MSc", "PhD">,
  "education_detail": "Full degree, specialisation, institution and year if available",
  "education_year": <4-digit year highest degree was completed, or null>,
  "jd_required_education": <EXACTLY one of the same education level strings if the JD specifies a minimum education requirement, else null>,
  "degree_match": <STRICT RULE — compare education_level against jd_required_education using this ordered hierarchy (lowest to highest): "10th" < "12th" < "Diploma" < "BSc/BA/BCom" < "BE/BTech" < "PG/MBA/MSc" < "PhD". Set FALSE if candidate education_level ranks LOWER than jd_required_education in that hierarchy. Set TRUE if equal or higher. Set null ONLY when jd_required_education is null.>,
  "degree_mismatch_note": <if degree_match is false, write a clear one-line explanation such as "JD requires BE/BTech but candidate holds a Diploma", else null>,
  "job_history": [
    {
      "company": "Company name",
      "role": "Job title at that company",
      "start": "YYYY-MM or YYYY",
      "end": "YYYY-MM, YYYY, or present",
      "tenure_months": <calculated integer; use current month/year if end is present>,
      "is_job_hop": <MUST be true if tenure_months < 12, false otherwise — no exceptions>
    }
  ],
  "job_hopping_flags": ["Company Name (Xm)" for every entry where is_job_hop is true],
  "career_gaps": [
    {
      "from": "YYYY-MM end of previous role",
      "to": "YYYY-MM start of next role",
      "gap_months": <integer>,
      "context": "Any explanation in resume (sabbatical, study, family etc.) or null"
    }
  ],
  "matched_skills": ["up to 8 skills in resume that match JD requirements"],
  "missing_skills": ["up to 6 key JD skills absent from resume"],
  "strengths": ["3 specific evidence-based strengths relevant to this role"],
  "concerns": ["up to 3 specific concerns or red flags"],
  "summary": "2-3 sentence overall evaluation for this specific role",
  "company_fit": {
    "candidate_current_company": "Name of candidate's most recent employer from job history",
    "candidate_company_employees": <integer — estimate global headcount; use actual data for well-known companies; for lesser-known, infer from resume clues; null only if absolutely no clues exist>,
    "candidate_company_size_estimate": "Always-present human-readable size string — e.g. '~500 employees', '~2,000 employees', '~50 employees — estimated'. Use actual headcount for well-known companies. For lesser-known, estimate from resume clues: team sizes mentioned, number of office locations, enterprise client names, regulatory audits (USFDA, ISO 9001 etc.), or industry norms. Append ' — estimated' when inferred. Never null.",
    "candidate_company_tier": <1 if 500+ employees | 2 if 200-499 | 3 if under 200. Always estimate using resume clues — null only if the resume provides no usable signal whatsoever>,
    "candidate_company_industry": "Industry sector of candidate's most recent employer — always provide; append ' — estimated' if inferred rather than known",
    "client_company_employees": <integer — estimate of hiring company headcount from your training data, or null if not specified or unknown>,
    "client_company_tier": <1, 2, or 3 using same thresholds | null if no client company specified or unknown>,
    "client_company_industry": "Industry of the hiring company, or null if not specified",
    "tier_gap": <abs(candidate_company_tier - client_company_tier) as integer | null if either tier is null>,
    "fit_direction": <"upgrade" if candidate moves from smaller to larger org (candidate tier# > client tier#) | "downgrade" if larger to smaller | "lateral" if same tier | null if unknown>,
    "mismatch_flag": <null if tier_gap is 0 or 1 or null — otherwise a concise one-sentence recruiter-facing risk note; always flag conservatively, never auto-reject>,
    "company_fit_score": <integer 0-10: rate alignment across all 5 dimensions below; 8-10=strong, 5-7=partial, 3-4=weak, 0-2=mismatch>,
    "company_fit_verdict": <EXACTLY one of "strong_fit" | "partial_fit" | "weak_fit" | "mismatch">,
    "company_fit_explanation": "3-4 sentences explaining the fit or mismatch across the dimensions that matter most for this JD and client",
    "company_size_gap_warning": <if candidate_company_tier number is HIGHER than client_company_tier number (i.e. candidate's org is smaller), provide one recruiter-facing warning line; otherwise null. Examples: candidate Tier 3 + client Tier 1 → "Candidate comes from a small company — may need time to adapt to a large enterprise environment"; candidate Tier 2 + client Tier 1 → "Candidate comes from a mid-size company — verify comfort with large enterprise processes and scale"; candidate Tier 3 + client Tier 2 → "Candidate comes from a smaller company — confirm readiness for mid-size organisation structure". Set null when candidate tier equals or is numerically lower than client tier (same size or larger org), or when either tier is unknown.>
  },
  "client_conflict": {
    "is_current_employee": <true if the candidate's current/most recent employer matches the CLIENT/HIRING COMPANY above using fuzzy case-insensitive partial matching — e.g. "Volkswagen" matches "Volkswagen India", "VW Group", "VWGTS", "Volkswagen AG". Set false if no client company was specified or no match found>,
    "is_ex_employee": <true ONLY if a PREVIOUS employer (not the most recent/current one) matches the client company using same fuzzy matching AND is_current_employee is false. Set false otherwise or if no client company specified>,
    "matched_company_name": "Exact company name string from the resume that matched, or null if no match",
    "ex_employee_period": <if is_ex_employee is true: { "company": "matched company name", "role": "role title at that company", "start": "YYYY-MM or YYYY", "end": "YYYY-MM or YYYY", "tenure_months": <integer> } — set to null if is_ex_employee is false>
  },
  "candidate_location": "City or region extracted from resume contact details, address, or LinkedIn city — null if not found anywhere in resume",
  "location_flag": <"green" if location matches required or role is remote/Pan-India | "amber" if neighbouring city pair or location not found in resume | "red" if different city/region requiring relocation | null if no required_location was provided>,
  "location_reason": "One-line recruiter-facing explanation — e.g. 'Remote role — location not a barrier', 'Location match — Pune', 'Based in Mumbai, role in Pune — likely open to relocation, confirm', 'Currently in Chennai, role requires Delhi — confirm relocation intent', 'Location not mentioned in resume — verify with candidate' — null when location_flag is null",
  "recommendation": <EXACTLY one of "proceed" | "hold" | "do_not_proceed" — your explicit hiring recommendation for this candidate>,
  "recommendation_reason": "Exactly 2 sentences, specific to this candidate and JD, explaining why you made this recommendation — not generic",
  "candidate_phone": "<10-digit Indian mobile number extracted from the resume — strip country code (+91/91/0), spaces, dashes; take the first mobile number found; return null if no phone number is present>"
}"""

_SCORING = (
    "Scoring guide:\n"
    "90-100 = Exceptional — exceeds all requirements\n"
    "75-89  = Strong — meets most requirements\n"
    "60-74  = Moderate — meets core requirements, minor gaps\n"
    "40-59  = Partial — missing key requirements or multiple concerns\n"
    "0-39   = Weak — degree mismatch, excessive job hopping, or major skill gaps\n"
    "Deductions: education below JD requirement −10 to −20 pts | "
    "each role with tenure < 12 months −5 pts | unexplained gap > 6 months −5 pts"
)

_JD_GAPS_SCHEMA = """{
  "gaps": [
    {
      "field": "experience_range|location|work_type|salary|tech_stack|industry|team_size|reporting",
      "severity": "high|medium|low",
      "note": "What is missing and why it matters for screening decisions"
    }
  ]
}"""

_JD_EXTRACT_SCHEMA = """{
  "extracted": {
    "min_experience": <integer years clearly stated in JD, or null>,
    "max_experience": <integer years clearly stated in JD, or null>,
    "must_have_skills": ["skill1", "skill2"],
    "nice_to_have_skills": ["skill1"],
    "client_company": "company name if mentioned in JD, else null",
    "work_location": "city or region (e.g. Pune, Bangalore, Mumbai) and/or work model (remote/hybrid/onsite) — capture whatever location info is stated; null if neither a city/region nor a work model is mentioned anywhere in the JD",
    "industry": "specific industry or domain if mentioned, else null",
    "product_client_check": "any specific product, platform, or company the candidate must have worked on OR must NOT have worked at (e.g. 'must have SAP S/4HANA experience', 'must not be from Tata Motors') — null if no such requirement is mentioned in the JD"
  },
  "missing_fields": ["field_name", ...]
}
Only include in missing_fields what is genuinely absent or too vague for screening.
Include work_location in missing_fields if the JD does not mention any city, region, or work model (remote/hybrid/onsite).
Include industry in missing_fields if the JD does not mention a preferred industry or domain background.
Include product_client_check in missing_fields if the JD does not specify any product/platform the candidate must know, or any company they must or must not have worked at (beyond what client_company captures).
Valid field_name values: min_experience, max_experience, must_have_skills, nice_to_have_skills, client_company, work_location, industry, product_client_check"""

_KEYWORD_SUGGEST_SCHEMA = """{
  "must_have_keywords": ["keyword1", "keyword2"],
  "nice_to_have_keywords": ["keyword1", "keyword2"]
}"""


# ─── Experience formatting helpers ───────────────────────────────────────────

def _fmt_exp_display(total_months: int | None) -> str:
    if not total_months:
        return "N/A"
    years, months = divmod(total_months, 12)
    if years == 0:
        return f"{months} mo"
    if months == 0:
        return f"{years} yr" if years == 1 else f"{years} yrs"
    return f"{years} yr {months} mo" if years == 1 else f"{years} yrs {months} mo"


def _fmt_exp_years(total_months: int | None) -> float:
    if not total_months:
        return 0.0
    return round(total_months / 12, 1)


# ─── JD Analysis (gaps panel) ─────────────────────────────────────────────────

def analyze_jd(job_description: str) -> list:
    """Identify critical recruiting information missing or too vague in the JD."""
    prompt = (
        "Analyze this job description. List critical screening information that is MISSING or too vague.\n\n"
        f"JOB DESCRIPTION:\n{job_description}\n\n"
        "Check these fields:\n"
        "- experience_range: minimum/maximum years of experience required\n"
        "- location: city, country, or remote/hybrid policy\n"
        "- work_type: remote / hybrid / onsite\n"
        "- salary: compensation range or band\n"
        "- tech_stack: specific technologies required (for technical roles)\n"
        "- industry: preferred industry background\n"
        "- team_size: size of team or company\n"
        "- reporting: who the role reports to\n\n"
        "Only flag fields that are genuinely absent. Do not flag fields that are mentioned.\n"
        f"Return ONLY this JSON:\n{_JD_GAPS_SCHEMA}"
    )
    try:
        response = client.chat.completions.create(
            model="llama-3.3-70b-versatile",
            max_tokens=600,
            temperature=0.1,
            messages=[
                {"role": "system", "content": "You identify missing information in job descriptions. Respond with valid JSON only."},
                {"role": "user", "content": prompt},
            ],
        )
        raw = response.choices[0].message.content.strip()
        raw = re.sub(r"^```(?:json)?\s*", "", raw, flags=re.MULTILINE)
        raw = re.sub(r"\s*```$", "", raw, flags=re.MULTILINE)
        return json.loads(raw).get("gaps", [])
    except Exception:
        return []


# ─── JD Field Extraction (Step 1 → 2) ────────────────────────────────────────

def extract_jd_context(job_description: str) -> dict:
    prompt = (
        "Analyze this job description and extract all screening-relevant information that is clearly stated.\n"
        "Then list fields that are genuinely missing or too vague for candidate screening.\n\n"
        f"JOB DESCRIPTION:\n{job_description}\n\n"
        f"Return ONLY this JSON:\n{_JD_EXTRACT_SCHEMA}"
    )
    try:
        response = client.chat.completions.create(
            model="llama-3.3-70b-versatile",
            max_tokens=700,
            temperature=0.1,
            messages=[
                {"role": "system", "content": "Extract structured information from job descriptions. Respond with valid JSON only."},
                {"role": "user", "content": prompt},
            ],
        )
        raw = response.choices[0].message.content.strip()
        raw = re.sub(r"^```(?:json)?\s*", "", raw, flags=re.MULTILINE)
        raw = re.sub(r"\s*```$", "", raw, flags=re.MULTILINE)
        return json.loads(raw)
    except Exception:
        return {
            "extracted": {
                "min_experience": None, "max_experience": None,
                "must_have_skills": [], "nice_to_have_skills": [],
                "client_company": None, "work_location": None, "industry": None,
                "product_client_check": None,
            },
            "missing_fields": ["min_experience", "max_experience", "must_have_skills", "client_company", "work_location", "product_client_check"],
        }


# ─── Keyword Suggestions (Step 2 → 3) ────────────────────────────────────────

def suggest_keywords_for_jd(job_description: str, context: dict) -> dict:
    parts = []
    if context.get("min_experience") or context.get("max_experience"):
        parts.append(f"- Experience: {context.get('min_experience', '?')}–{context.get('max_experience', '?')} years")
    if context.get("must_have_skills"):
        parts.append(f"- Must-have skills already listed: {context['must_have_skills']}")
    if context.get("nice_to_have_skills"):
        parts.append(f"- Nice-to-have already listed: {context['nice_to_have_skills']}")
    if context.get("industry"):
        parts.append(f"- Industry: {context['industry']}")
    if context.get("work_location"):
        parts.append(f"- Work type: {context['work_location']}")

    context_str = "\n".join(parts) if parts else "No additional context provided."

    prompt = (
        "You are a senior technical recruiter with deep market knowledge. "
        "Suggest additional relevant keywords for this role that are NOT already present in the JD or context below.\n\n"
        f"JOB DESCRIPTION:\n{job_description}\n\n"
        f"ROLE CONTEXT:\n{context_str}\n\n"
        "Must-have keywords: industry-standard skills, tools, certifications ESSENTIAL for this role.\n"
        "Nice-to-have keywords: beneficial but not critical — adjacent skills, complementary tools.\n"
        "Return up to 12 per category. Each keyword should be 1–4 words max.\n"
        f"Return ONLY this JSON:\n{_KEYWORD_SUGGEST_SCHEMA}"
    )
    try:
        response = client.chat.completions.create(
            model="llama-3.3-70b-versatile",
            max_tokens=500,
            temperature=0.4,
            messages=[
                {"role": "system", "content": "You suggest relevant job keywords. Respond with valid JSON only."},
                {"role": "user", "content": prompt},
            ],
        )
        raw = response.choices[0].message.content.strip()
        raw = re.sub(r"^```(?:json)?\s*", "", raw, flags=re.MULTILINE)
        raw = re.sub(r"\s*```$", "", raw, flags=re.MULTILINE)
        return json.loads(raw)
    except Exception:
        return {"must_have_keywords": [], "nice_to_have_keywords": []}


# ─── Text Extractors ──────────────────────────────────────────────────────────

def extract_pdf_text(data: bytes) -> str:
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        pages = [p.extract_text() or "" for p in pdf.pages]
    return "\n\n".join(pages)


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def extract_docx_text(data: bytes) -> str:
    segments: list[str] = []

    try:
        doc = docx.Document(io.BytesIO(data))
        m1: list[str] = []

        def _drain_table(tbl) -> None:
            seen_tcs: set = set()
            for row in tbl.rows:
                cell_texts: list[str] = []
                for cell in row.cells:
                    tc_id = id(cell._tc)
                    if tc_id in seen_tcs:
                        continue
                    seen_tcs.add(tc_id)
                    ct = cell.text.strip()
                    if ct:
                        cell_texts.append(ct)
                    for nested in cell.tables:
                        _drain_table(nested)
                if cell_texts:
                    m1.append("  |  ".join(cell_texts))

        for p in doc.paragraphs:
            if p.text.strip():
                m1.append(p.text.strip())
        for tbl in doc.tables:
            _drain_table(tbl)
        for sec in doc.sections:
            for part in (
                sec.header, sec.footer,
                sec.first_page_header, sec.first_page_footer,
                sec.even_page_header, sec.even_page_footer,
            ):
                try:
                    for p in part.paragraphs:
                        if p.text.strip():
                            m1.append(p.text.strip())
                    for tbl in part.tables:
                        _drain_table(tbl)
                except Exception:
                    pass
        segments.append("\n".join(m1))
    except Exception:
        pass

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            targets = [
                n for n in z.namelist()
                if n.startswith("word/") and n.endswith(".xml")
                and any(k in n for k in ("document", "header", "footer"))
            ]
            m2: list[str] = []
            for name in targets:
                try:
                    root = ET.fromstring(z.read(name))
                    for para in root.iter(f"{_W}p"):
                        runs = [t.text for t in para.iter(f"{_W}t") if t.text]
                        text = "".join(runs).strip()
                        if text:
                            m2.append(text)
                except Exception:
                    pass
            segments.append("\n".join(m2))
    except Exception:
        pass

    try:
        import docx2txt
        txt = docx2txt.process(io.BytesIO(data))
        if txt and txt.strip():
            segments.append(txt.strip())
    except Exception:
        pass

    seen: set[str] = set()
    out: list[str] = []
    for block in segments:
        for raw_line in block.splitlines():
            line = raw_line.strip()
            norm = " ".join(line.lower().split())
            if norm and norm not in seen:
                seen.add(norm)
                out.append(line)

    return "\n".join(out)


# ─── LLM Scoring ─────────────────────────────────────────────────────────────

def call_llm(
    resume_text: str,
    job_description: str,
    filename: str,
    client_company: str = "",
    enriched_context: str = "",
    required_location: str = "",
    benchmark_fingerprint: str = "",
    client_learning_note: str = "",
) -> dict:
    client_note = (
        f"\nCLIENT/HIRING COMPANY: {client_company}\n"
        "Use your training knowledge to estimate this company's employee count, tier, and industry for the company_fit field.\n"
        "For client_conflict: check if the candidate currently works at (is_current_employee) or previously worked at (is_ex_employee) this company. "
        "Use fuzzy case-insensitive partial matching — e.g. 'Volkswagen' matches 'Volkswagen India', 'VW Group', 'VWGTS', 'Volkswagen AG'."
        if client_company.strip() else
        "\nCLIENT/HIRING COMPANY: Not specified — set client_company_employees, client_company_tier, client_company_industry to null, "
        "and set client_conflict.is_current_employee and client_conflict.is_ex_employee to false with matched_company_name and ex_employee_period as null."
    )
    enriched_note = (
        f"\n\nRECRUITER-SPECIFIED REQUIREMENTS (use these for scoring, skill matching, and relevance):\n{enriched_context}"
        if enriched_context.strip() else ""
    )
    if required_location.strip():
        location_note = (
            f"\n\nLOCATION ASSESSMENT — Required location: {required_location}\n"
            "Extract the candidate's current city/region from the resume (contact section, address, LinkedIn city, or any location reference). "
            "Set candidate_location to this value (or null if not found).\n\n"
            "STEP 1 — NORMALISE CITY NAMES: Before comparing, map both the candidate city and the required city to their "
            "canonical name using the aliases below. All aliases for the same entry are an exact match:\n"
            "• Chhatrapati Sambhajinagar: Aurangabad, Aurangabad (MH), Ch. Sambhajinagar, Chhatrapati Sambhajinagar, "
            "Chattrapati Sambhajinagar, CSN, C.S. Nagar, Sambhajinagar, CSNG, Chh. Sambhajinagar\n"
            "• Ahilyanagar: Ahmednagar, Ahmed Nagar, Ahmedanagar, Ahilya Nagar, Ahilyanagar\n"
            "• Mumbai: Bombay, Mumbai, Navi Mumbai, Thane, Mira Road, Vasai, Virar, Kalyan, Dombivli, MMRDA region\n"
            "• Pune: Pune, Poona, Pimpri, Chinchwad, Pimpri-Chinchwad, PCMC, Hinjewadi\n"
            "• Delhi: Delhi, New Delhi, NCR, Delhi NCR, National Capital Region\n"
            "• Bengaluru: Bangalore, Bengaluru, Banglore, Bengaluru Urban\n"
            "• Chennai: Chennai, Madras\n"
            "• Kolkata: Kolkata, Calcutta\n"
            "• Vadodara: Vadodara, Baroda\n"
            "• Gurugram: Gurgaon, Gurugram, Gurugraam\n"
            "• Noida: Noida, Greater Noida, Gr. Noida\n"
            "• Nashik: Nashik, Nasik\n"
            "• Mysuru: Mysore, Mysuru\n"
            "• Thiruvananthapuram: Trivandrum, Thiruvananthapuram\n"
            "• Kochi: Cochin, Kochi, Ernakulam\n"
            "• Belagavi: Belgaum, Belagavi\n"
            "• Hubballi: Hubli, Hubballi, Hubli-Dharwad\n"
            "• Shivamogga: Shimoga, Shivamogga\n"
            "• Mangaluru: Mangalore, Mangaluru\n"
            "• Vijayawada: Vijayawada, Bezawada\n"
            "• Visakhapatnam: Vizag, Visakhapatnam, Vishakhapatnam\n"
            "• Tiruchirapalli: Trichy, Tiruchirappalli, Tiruchirapalli\n"
            "• Kozhikode: Calicut, Kozhikode\n\n"
            "STEP 2 — APPLY LOCATION FLAG RULES (check in this exact order, stop at the first match):\n"
            "1. required_location contains 'remote', 'WFH', 'work from home', 'Pan-India', or 'anywhere' → "
            "green, 'Remote role — location not a barrier'\n"
            "2. After normalisation, candidate city == required city → green, 'Location match — [city]'\n"
            "3. Canonical cities form a known neighbouring pair — "
            "Pune–Mumbai, Pune–Chhatrapati Sambhajinagar, Delhi–Gurugram, Delhi–Noida, Delhi–Faridabad, "
            "Bengaluru–Mysuru, Hyderabad–Secunderabad, Chennai–Tambaram → "
            "amber, 'Based in [x], role in [y] — likely open to relocation, confirm'\n"
            "4. Candidate location not found anywhere in resume → "
            "amber, 'Location not mentioned in resume — verify with candidate'\n"
            "5. City not in alias list above but inferred to be in the same state as the required city → "
            "amber, 'Based in [x] ([state]), role in [y] — different city, same state, confirm'\n"
            "6. Different state or distant city → red, 'Currently in [x], role requires [y] — confirm relocation intent'"
        )
    else:
        location_note = (
            "\n\nLOCATION: No required location provided — extract candidate_location from resume if available; "
            "set location_flag and location_reason to null."
        )
    company_fit_note = (
        "\n\nCOMPANY FIT ANALYSIS — Score the candidate across these 5 dimensions vs the client/JD:\n"
        "1. INDUSTRY ALIGNMENT: Same or adjacent sector? (automotive→automotive ancillary = strong; BFSI→retail = weak)\n"
        "2. PRODUCT & PLATFORM RELEVANCE: Has the candidate worked on same product/platform/stack the JD requires? "
        "(SAP S/4HANA, Salesforce, AWS etc.) — more important than company size.\n"
        "3. PROCESS & DOMAIN DEPTH: Exposed to similar business processes/workflows? "
        "(supply chain for mfg, P&L for finance, agile delivery for tech)\n"
        "4. COMPANY SIZE & COMPLEXITY: Use existing Tier 1/2/3 as ONE input — a Tier 3 company with deep domain "
        "can outrank a Tier 1 with no domain relevance.\n"
        "5. Use your training knowledge for well-known companies. For lesser-known companies, infer from the work done, "
        "technologies, clients, and projects mentioned in the resume. Never say 'I don't know this company' — always make a reasoned assessment.\n"
        "CANDIDATE COMPANY ESTIMATION — Always fill candidate_company_tier, candidate_company_industry, and "
        "candidate_company_size_estimate even for obscure or small companies. Use any available resume signal: "
        "explicit team/headcount mentions, number of office locations, names of enterprise clients (implying vendor scale), "
        "regulatory audits passed (USFDA, ISO 9001, IATF etc.), or industry norms for that category of firm. "
        "Append ' — estimated' to candidate_company_size_estimate when the value is inferred. "
        "Leave null only when the resume provides literally no usable signal.\n"
        "company_fit_verdict: \"strong_fit\" if 3+ dimensions align well; \"partial_fit\" if 1-2 align; "
        "\"weak_fit\" if very little alignment; \"mismatch\" if clear mismatch likely to cause client rejection.\n"
        "COMPANY SIZE GAP WARNING RULE: Set company_size_gap_warning whenever candidate_company_tier number > client_company_tier number "
        "(meaning the candidate's company is smaller). This warning is independent of the overall verdict — "
        "a strong domain fit does NOT cancel the size gap warning. "
        "Tier 3 candidate + Tier 1 client → 'Candidate comes from a small company — may need time to adapt to a large enterprise environment'. "
        "Tier 2 candidate + Tier 1 client → 'Candidate comes from a mid-size company — verify comfort with large enterprise processes and scale'. "
        "Tier 3 candidate + Tier 2 client → 'Candidate comes from a smaller company — confirm readiness for mid-size organisation structure'. "
        "Set null when candidate tier equals or is numerically lower than client tier, or when either tier is unknown."
    )
    recommendation_rules = (
        "\n\nRECOMMENDATION RULES (apply after full analysis):\n"
        "- \"proceed\": score ≥ 75 AND no hard red flags (active client conflict, degree below requirement, critical missing must-have skill)\n"
        "- \"hold\": score 50–74 OR one significant concern that needs verification before a decision\n"
        "- \"do_not_proceed\": score < 50 OR hard disqualifier present (client conflict, completely wrong domain, "
        "non-negotiable missing skill explicitly stated in JD)\n"
        "- These are guidelines — use judgement on the full picture. A score-75 candidate with an obvious red flag "
        "should be \"hold\" or \"do_not_proceed\". A score-48 candidate with a rare niche skill the JD specifically calls out can be \"hold\".\n"
        "- recommendation_reason must name the specific skill/flag/score driver — never write generic phrases like \"overall fit is good\"."
    )
    benchmark_note = ""
    if benchmark_fingerprint.strip():
        benchmark_note = (
            f"\n\nBENCHMARK CANDIDATE FINGERPRINT (IDEAL PROFILE):\n{benchmark_fingerprint}\n\n"
            "Compare this candidate against the above ideal profile. "
            "Add these extra fields to your JSON response (do NOT change the existing score field):\n"
            "  \"benchmark_match_percent\": integer 0-100 (how well this candidate matches the ideal profile),\n"
            "  \"benchmark_similarities\": [\"string describing a specific area where the candidate matches the benchmark\"],\n"
            "  \"benchmark_differences\": [\"string describing a specific gap vs the benchmark\"],\n"
            "  \"benchmark_verdict\": one of \"strong_match\" | \"good_match\" | \"partial_match\" | \"weak_match\"\n"
            "Score benchmark_match_percent by comparing: skills overlap (30%), experience level (20%), "
            "domain match (25%), education level (10%), company tier and career stability (15%)."
        )
    learning_note = (
        f"\n\nCLIENT-SPECIFIC INTELLIGENCE (based on historical outcomes for {client_company}):\n"
        f"{client_learning_note}\n"
        "Apply extra weight to these factors when scoring this candidate."
        if client_learning_note.strip() else ""
    )
    prompt = (
        "Analyze this resume against the job description and return a JSON object.\n\n"
        f"JOB DESCRIPTION:\n{job_description}{client_note}{enriched_note}{location_note}{company_fit_note}{benchmark_note}{learning_note}\n\n---\n\n"
        f"RESUME:\n{resume_text[:MAX_RESUME_CHARS]}\n\n---\n\n"
        f"Return ONLY this JSON structure (no markdown, no code blocks):\n{_JSON_SCHEMA}\n\n"
        f"Scoring:\n{_SCORING}{recommendation_rules}"
    )
    response = None
    for attempt in range(4):  # up to 3 retries after the first attempt
        try:
            response = client.chat.completions.create(
                model="llama-3.3-70b-versatile",
                max_tokens=3000,
                temperature=0.1,
                messages=[
                    {"role": "system", "content": SYSTEM_PROMPT},
                    {"role": "user", "content": prompt},
                ],
            )
            break  # success
        except groq.RateLimitError as exc:
            if attempt < 3:
                time.sleep(10)  # wait 10 s then retry
            else:
                raise  # exhausted all retries

    raw = response.choices[0].message.content.strip()
    raw = re.sub(r"^```(?:json)?\s*", "", raw, flags=re.MULTILINE)
    raw = re.sub(r"\s*```$", "", raw, flags=re.MULTILINE)
    result = json.loads(raw)

    placeholder = result.get("candidate_name", "")
    if not placeholder or "full name" in placeholder.lower():
        stem = filename.rsplit(".", 1)[0].replace("_", " ").replace("-", " ")
        result["candidate_name"] = " ".join(w.capitalize() for w in stem.split())

    # Override AI-computed display fields with precise Python calculations
    total_m    = result.get("experience_total_months")
    relevant_m = result.get("experience_relevant_months")
    result["total_experience_years"]      = _fmt_exp_years(total_m)
    result["total_experience_display"]    = _fmt_exp_display(total_m)
    result["relevant_experience_years"]   = _fmt_exp_years(relevant_m)
    result["relevant_experience_display"] = _fmt_exp_display(relevant_m)

    return result


async def process_resume(
    file_bytes: bytes,
    filename: str,
    idx: int,
    jd: str,
    client_company: str = "",
    enriched_context: str = "",
    required_location: str = "",
    benchmark_fingerprint: str = "",
    client_learning_note: str = "",
):
    async with _semaphore:
        try:
            ext = filename.lower().rsplit(".", 1)[-1]
            if ext == "pdf":
                text = await asyncio.to_thread(extract_pdf_text, file_bytes)
            elif ext in ("docx", "doc"):
                text = await asyncio.to_thread(extract_docx_text, file_bytes)
            else:
                return None, {"filename": filename, "error": "Unsupported format (use PDF or DOCX)"}

            if not text.strip():
                return None, {"filename": filename, "error": "No text could be extracted (scanned PDF?)"}

            result = await asyncio.to_thread(call_llm, text, jd, filename, client_company, enriched_context, required_location, benchmark_fingerprint, client_learning_note)
            result.update({"id": idx, "filename": filename, "status": "pending"})
            return result, None

        except json.JSONDecodeError:
            return None, {"filename": filename, "error": "AI returned unparseable response"}
        except Exception as exc:
            return None, {"filename": filename, "error": str(exc)}


# ─── Endpoints ────────────────────────────────────────────────────────────────

@app.post("/api/analyze-jd")
async def analyze_jd_endpoint(
    job_description: str = Form(""),
    jd_file: Optional[UploadFile] = File(None),
):
    """Step 1 → 2: Extract fields from JD and identify what is missing."""
    if jd_file and jd_file.filename:
        file_bytes = await jd_file.read()
        ext = (jd_file.filename or "").lower().rsplit(".", 1)[-1]
        if ext == "pdf":
            job_description = await asyncio.to_thread(extract_pdf_text, file_bytes)
        elif ext in ("docx", "doc"):
            job_description = await asyncio.to_thread(extract_docx_text, file_bytes)
        else:
            raise HTTPException(status_code=400, detail="JD file must be PDF or DOCX")

    if not job_description.strip():
        raise HTTPException(status_code=400, detail="Job description text or file is required")

    result = await asyncio.to_thread(extract_jd_context, job_description)
    return {"job_description": job_description, **result}


@app.post("/api/suggest-keywords")
async def suggest_keywords_endpoint(
    job_description: str = Form(...),
    min_experience: str = Form(""),
    max_experience: str = Form(""),
    must_have_skills: str = Form(""),
    nice_to_have_skills: str = Form(""),
    client_company: str = Form(""),
    work_location: str = Form(""),
    industry: str = Form(""),
):
    """Step 2 → 3: Suggest additional keywords missing from the JD."""
    context = {k: v for k, v in {
        "min_experience": min_experience,
        "max_experience": max_experience,
        "must_have_skills": must_have_skills,
        "nice_to_have_skills": nice_to_have_skills,
        "client_company": client_company,
        "work_location": work_location,
        "industry": industry,
    }.items() if v.strip()}
    result = await asyncio.to_thread(suggest_keywords_for_jd, job_description, context)
    return result


@app.post("/api/projects/{project_id}/benchmark")
async def upload_benchmark_cv(project_id: int, benchmark_cv: UploadFile = File(...)):
    file_bytes = await benchmark_cv.read()
    fn = (benchmark_cv.filename or "").lower()
    if fn.endswith(".pdf"):
        cv_text = await asyncio.to_thread(extract_pdf_text, file_bytes)
    elif fn.endswith((".docx", ".doc")):
        cv_text = await asyncio.to_thread(extract_docx_text, file_bytes)
    else:
        raise HTTPException(status_code=400, detail="Only PDF and DOCX files are accepted")

    if not cv_text.strip():
        raise HTTPException(status_code=400, detail="No text could be extracted from the benchmark CV")

    fingerprint_prompt = f"""You are an expert recruiter. Analyse this CV and extract the ideal candidate fingerprint.
Return ONLY valid JSON with exactly these fields:
{{
  "candidate_name": "string",
  "current_role": "string",
  "total_experience_years": number,
  "key_skills": ["string"],
  "domain_expertise": ["string"],
  "education_level": "string",
  "company_tier": 1,
  "career_stability": "stable",
  "standout_qualities": ["string"],
  "summary": "string"
}}
company_tier must be 1, 2, or 3 (integer).
career_stability must be "stable", "average", or "unstable".

CV TEXT:
{cv_text[:6000]}"""

    fp_response = client.chat.completions.create(
        model="llama-3.3-70b-versatile",
        messages=[
            {"role": "system", "content": "You are an expert recruiter. Respond with valid JSON only — no markdown, no code fences."},
            {"role": "user", "content": fingerprint_prompt},
        ],
        temperature=0.1,
        max_tokens=800,
    )
    raw = fp_response.choices[0].message.content.strip()
    raw = re.sub(r"^```(?:json)?\s*", "", raw, flags=re.MULTILINE)
    raw = re.sub(r"\s*```$", "", raw, flags=re.MULTILINE)
    fingerprint = json.loads(raw)

    db = SessionLocal()
    try:
        project = db.query(DBProject).filter(DBProject.id == project_id).first()
        if not project:
            raise HTTPException(status_code=404, detail="Project not found")
        project.benchmark_fingerprint = json.dumps(fingerprint)
        db.commit()
    finally:
        db.close()

    return fingerprint


def _compute_benchmark_for_candidate(result_dict: dict, fingerprint: dict) -> dict:
    candidate_summary = (
        f"Candidate: {result_dict.get('candidate_name', 'Unknown')}\n"
        f"Current Role: {result_dict.get('current_role', 'N/A')}\n"
        f"Total Experience: {result_dict.get('total_experience_display', 'N/A')} ({result_dict.get('total_experience_years', 0)} yrs)\n"
        f"Relevant Experience: {result_dict.get('relevant_experience_display', 'N/A')}\n"
        f"Education: {result_dict.get('education_level', 'N/A')} — {result_dict.get('education_detail', '')}\n"
        f"Matched Skills: {', '.join(result_dict.get('matched_skills', []))}\n"
        f"Missing Skills: {', '.join(result_dict.get('missing_skills', []))}\n"
        f"Company Tier: {(result_dict.get('company_fit') or {}).get('candidate_company_tier', 'N/A')}\n"
        f"Company Industry: {(result_dict.get('company_fit') or {}).get('candidate_company_industry', 'N/A')}\n"
        f"Job Hopping Flags: {', '.join(result_dict.get('job_hopping_flags', [])) or 'None'}\n"
        f"Strengths: {'; '.join(result_dict.get('strengths', []))}\n"
        f"Score: {result_dict.get('score', 'N/A')}/100"
    )
    fp_str = json.dumps(fingerprint, indent=2)
    prompt = (
        f"Compare this candidate against the benchmark ideal profile and return a JSON object.\n\n"
        f"BENCHMARK IDEAL PROFILE:\n{fp_str}\n\n"
        f"CANDIDATE PROFILE:\n{candidate_summary}\n\n"
        f"Evaluate match across: skills overlap (30%), experience level (20%), domain match (25%), "
        f"education level (10%), company tier and career stability (15%).\n\n"
        f"Return ONLY this JSON (no markdown, no code blocks):\n"
        f'{{\n'
        f'  "benchmark_match_percent": <integer 0-100>,\n'
        f'  "benchmark_similarities": ["specific area where candidate matches the benchmark"],\n'
        f'  "benchmark_differences": ["specific gap vs the benchmark"],\n'
        f'  "benchmark_verdict": <exactly one of "strong_match" | "good_match" | "partial_match" | "weak_match">\n'
        f'}}'
    )
    response = client.chat.completions.create(
        model="llama-3.3-70b-versatile",
        max_tokens=600,
        temperature=0.1,
        messages=[
            {"role": "system", "content": "You compare candidates against benchmark profiles. Respond with valid JSON only."},
            {"role": "user", "content": prompt},
        ],
    )
    raw = response.choices[0].message.content.strip()
    raw = re.sub(r"^```(?:json)?\s*", "", raw, flags=re.MULTILINE)
    raw = re.sub(r"\s*```$", "", raw, flags=re.MULTILINE)
    return json.loads(raw)


@app.post("/api/projects/{project_id}/apply-benchmark")
async def apply_benchmark(project_id: int):
    db = SessionLocal()
    try:
        project = db.query(DBProject).filter(DBProject.id == project_id).first()
        if not project:
            raise HTTPException(status_code=404, detail="Project not found")
        if not project.benchmark_fingerprint:
            raise HTTPException(status_code=400, detail="No benchmark fingerprint set for this project")

        fingerprint = json.loads(project.benchmark_fingerprint)
        candidates = (
            db.query(DBCandidate)
            .filter(DBCandidate.project_id == project_id)
            .all()
        )
        if not candidates:
            return {"updated": 0, "candidates": []}

        updated_candidates = []
        for i, candidate in enumerate(candidates):
            if i > 0:
                await asyncio.sleep(8)
            try:
                result_dict = json.loads(candidate.result_json or "{}")
                bench_data = await asyncio.to_thread(_compute_benchmark_for_candidate, result_dict, fingerprint)
                result_dict["benchmark_match_percent"] = bench_data.get("benchmark_match_percent")
                result_dict["benchmark_similarities"] = bench_data.get("benchmark_similarities", [])
                result_dict["benchmark_differences"] = bench_data.get("benchmark_differences", [])
                result_dict["benchmark_verdict"] = bench_data.get("benchmark_verdict")
                candidate.result_json = json.dumps(result_dict)
                candidate.benchmark_match_percent = bench_data.get("benchmark_match_percent")
                candidate.benchmark_similarities = json.dumps(bench_data.get("benchmark_similarities", []))
                candidate.benchmark_differences = json.dumps(bench_data.get("benchmark_differences", []))
                candidate.benchmark_verdict = bench_data.get("benchmark_verdict")
            except Exception as exc:
                print(f"[apply-benchmark] candidate {candidate.id} failed: {exc}")

        db.commit()

        for candidate in candidates:
            updated_candidates.append({
                "id": candidate.id,
                "name": candidate.name,
                "filename": candidate.filename,
                "score": candidate.score,
                "result_json": candidate.result_json,
                "status": candidate.status or "pending",
                "remark": candidate.remark,
                "linkedin_url": candidate.linkedin_url,
                "linkedin_text": candidate.linkedin_text,
                "uploaded_at": candidate.uploaded_at.isoformat(),
            })

        return {"updated": len(candidates), "candidates": updated_candidates}
    finally:
        db.close()


@app.post("/api/analyze")
async def analyze_resumes(
    job_description: str = Form(...),
    client_company: str = Form(""),
    min_experience: str = Form(""),
    max_experience: str = Form(""),
    must_have_skills: str = Form(""),
    nice_to_have_skills: str = Form(""),
    work_location: str = Form(""),
    industry: str = Form(""),
    product_client_check: str = Form(""),
    benchmark_fingerprint: str = Form(""),
    resumes: List[UploadFile] = File(...),
):
    if not job_description.strip():
        raise HTTPException(status_code=400, detail="Job description is required")
    if not resumes:
        raise HTTPException(status_code=400, detail="At least one resume is required")
    if len(resumes) > 5:
        raise HTTPException(status_code=400, detail="Maximum 5 resumes per batch")

    # Build enriched context string to improve scoring accuracy
    enriched_parts = []
    if min_experience or max_experience:
        enriched_parts.append(f"Experience Required: {min_experience or '?'}–{max_experience or '?'} years")
    if must_have_skills:
        enriched_parts.append(f"Must-have Skills: {must_have_skills}")
    if nice_to_have_skills:
        enriched_parts.append(f"Nice-to-have Skills: {nice_to_have_skills}")
    if work_location:
        enriched_parts.append(f"Work Location: {work_location}")
    if industry:
        enriched_parts.append(f"Industry: {industry}")
    if product_client_check:
        enriched_parts.append(f"Product/Platform/Client Constraints: {product_client_check}")
    enriched_context = "\n".join(enriched_parts)

    file_data = []
    for i, f in enumerate(resumes, start=1):
        file_data.append((await f.read(), f.filename or f"resume_{i}", i))

    jd_gaps = await asyncio.to_thread(analyze_jd, job_description)

    # Build client learning note from established rules (confidence >= 3)
    client_learning_note = ""
    if client_company.strip() and client_company.strip().lower() != "unknown":
        db = SessionLocal()
        try:
            rules = db.query(ClientLearning).filter(
                ClientLearning.client_name == client_company.strip(),
                ClientLearning.confidence >= 3,
            ).all()
            if rules:
                client_learning_note = "\n".join(f"- {r.learning_text}" for r in rules)
        finally:
            db.close()

    results = []
    errors = []
    for i, (data, name, idx) in enumerate(file_data):
        if i > 0:
            await asyncio.sleep(8)
        result, error = await process_resume(data, name, idx, job_description, client_company, enriched_context, work_location, benchmark_fingerprint, client_learning_note)
        if result is not None:
            results.append(result)
        if error is not None:
            errors.append(error)

    results.sort(key=lambda x: x.get("score", 0), reverse=True)

    return {"candidates": results, "errors": errors, "total": len(results), "jd_gaps": jd_gaps}


# ─── Call Guide Endpoint ──────────────────────────────────────────────────────

@app.post("/api/call-guide")
async def generate_call_guide(req: CallGuideRequest):
    c = req.candidate_json
    candidate_name = c.get("candidate_name", "Candidate")
    summary_parts = [
        f"Candidate: {candidate_name}",
        f"Score: {c.get('score', '?')}/100",
        f"Recommendation: {c.get('recommendation', '')}",
        f"Total Experience: {c.get('total_experience_display', '')}",
        f"Relevant Experience: {c.get('relevant_experience_display', '')}",
        f"Current Role: {c.get('current_role', '')}",
        f"Matched Skills: {', '.join(c.get('matched_skills', []))}",
        f"Missing Skills: {', '.join(c.get('missing_skills', []))}",
        f"Strengths: {'; '.join(c.get('strengths', [])[:3])}",
        f"Concerns: {'; '.join(c.get('concerns', [])[:3])}",
    ]
    candidate_summary = "\n".join(summary_parts)

    system_msg = (
        "You are an expert technical recruiter. "
        "Generate interview questions as a JSON array. "
        "Respond with valid JSON only — no markdown, no code fences, no extra text."
    )
    user_msg = f"""Job Description:
{req.jd_text[:3000]}

Candidate Profile:
{candidate_summary}

Generate 6 to 8 technical interview questions for a recruiter screening call.
Rules:
- Questions must be specific to the JD and this candidate's background
- Mix: verify claimed skills, probe experience depth, uncover identified gaps
- Each question needs a brief "what_to_listen_for" (2-3 lines) in plain language for a non-technical recruiter
- Questions must be open-ended, not yes/no
- Mark exactly 2 questions as must_ask: true (the most critical for this role)

Return only a JSON array:
[
  {{
    "question": "string",
    "what_to_listen_for": "string",
    "must_ask": false
  }}
]"""

    response = client.chat.completions.create(
        model="llama-3.3-70b-versatile",
        messages=[
            {"role": "system", "content": system_msg},
            {"role": "user", "content": user_msg},
        ],
        temperature=0.4,
        max_tokens=2000,
    )
    raw = response.choices[0].message.content.strip()
    raw = re.sub(r"^```(?:json)?\s*", "", raw, flags=re.MULTILINE)
    raw = re.sub(r"\s*```$", "", raw, flags=re.MULTILINE)
    questions = json.loads(raw)
    return {"candidate_name": candidate_name, "technical_questions": questions}


# ─── WhatsApp Message Endpoint ────────────────────────────────────────────────

@app.post("/api/candidates/{candidate_id}/whatsapp-message")
async def generate_whatsapp_message(candidate_id: int, body: WhatsAppMessageRequest):
    db = SessionLocal()
    try:
        candidate = db.query(DBCandidate).filter(DBCandidate.id == candidate_id).first()
        if not candidate:
            raise HTTPException(status_code=404, detail="Candidate not found")
        project = db.query(DBProject).filter(DBProject.id == candidate.project_id).first()
        result = json.loads(candidate.result_json or "{}")
        first_name = (result.get("candidate_name") or candidate.name or "there").split()[0]
        job_title = ""
        industry = ""
        if project:
            jd_ctx = json.loads(project.jd_context or "{}")
            job_title = jd_ctx.get("job_title") or ""
            industry = jd_ctx.get("industry") or jd_ctx.get("client_industry") or ""

        if body.message_type == "check_interest":
            prompt = (
                f"You are a professional recruiter writing a WhatsApp message to a candidate. "
                f"Write a concise, warm, professional WhatsApp message to check if the candidate "
                f"is interested in a job opportunity.\n\n"
                f"The message should:\n"
                f"- Open with a warm greeting using the candidate's first name\n"
                f"- Briefly introduce yourself as a recruiter from Agile Technology Solutions\n"
                f"- Mention the role and company type (without revealing confidential client details)\n"
                f"- Ask if they are open to exploring this opportunity\n"
                f"- Keep it under 100 words\n"
                f"- Use a conversational WhatsApp tone — not too formal\n"
                f"- End with: {body.recruiter_name}\n\n"
                f"Candidate first name: {first_name}\n"
                f"Role: {job_title or 'a senior position'}\n"
                f"Industry/domain: {industry or 'technology'}\n\n"
                f"Return only the message text. No formatting, no labels."
            )
        else:
            prompt = (
                f"You are a professional recruiter writing a WhatsApp message to a candidate. "
                f"Write a concise, warm, professional WhatsApp message based on the context provided.\n\n"
                f"Context from recruiter: {body.context}\n"
                f"Candidate first name: {first_name}\n"
                f"Role: {job_title or 'a position'}\n\n"
                f"Rules:\n"
                f"- Under 120 words\n"
                f"- Conversational WhatsApp tone\n"
                f"- Professional but warm\n"
                f"- End with: {body.recruiter_name}\n\n"
                f"Return only the message text."
            )
        resp = client.chat.completions.create(
            model="llama-3.3-70b-versatile",
            max_tokens=250,
            temperature=0.85,
            messages=[{"role": "user", "content": prompt}],
        )
        message = resp.choices[0].message.content.strip()
        return {"message": message, "candidate_phone": candidate.candidate_phone}
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))
    finally:
        db.close()


# ─── Remark Endpoints ─────────────────────────────────────────────────────────

@app.post("/api/candidates/{candidate_id}/remark")
async def generate_remark(candidate_id: int, req: RemarkGenerate):
    if not req.linkedin_url.strip() or not req.linkedin_text.strip():
        raise HTTPException(
            status_code=400,
            detail="LinkedIn URL and profile text are required to generate a remark",
        )
    db = SessionLocal()
    try:
        candidate = db.query(DBCandidate).filter(DBCandidate.id == candidate_id).first()
        if not candidate:
            raise HTTPException(status_code=404, detail="Candidate not found")
        c = json.loads(candidate.result_json or '{}')
        candidate_name = c.get("candidate_name", "Candidate")
        summary_parts = [
            f"Candidate: {candidate_name}",
            f"Score: {c.get('score', '?')}/100",
            f"Recommendation: {c.get('recommendation', '')}",
            f"Current Role: {c.get('current_role', '')}",
            f"Total Experience: {c.get('total_experience_display', '')}",
            f"Relevant Experience: {c.get('relevant_experience_display', '')}",
            f"Matched Skills: {', '.join(c.get('matched_skills', []))}",
            f"Strengths: {'; '.join(c.get('strengths', [])[:4])}",
            f"AI Summary: {c.get('summary', '')}",
        ]
        candidate_summary = "\n".join(summary_parts)

        system_msg = (
            "You are a professional senior recruiter writing a candidate endorsement remark to share with a hiring manager. "
            "You have three sources of information: the AI screening analysis of the candidate's CV, the candidate's LinkedIn profile, "
            "and notes from a recruiter call with the candidate. "
            "Return only the remark text — no headings, no bullet points, plain paragraph."
        )
        user_msg = f"""Candidate being evaluated: {candidate_name} | Current Role: {c.get('current_role', '')} | Company: {c.get('current_company', '')}

CV Analysis:
{candidate_summary}

LinkedIn Profile URL: {req.linkedin_url}
LinkedIn Profile Text:
{req.linkedin_text[:4000]}

Recruiter Call Notes:
{req.voice_transcript or 'No call notes provided'}

Write a concise, professional recruiter remark of 5-6 lines.
Rules:
- Open with the candidate's full name, current designation, and company
- Highlight 2-3 strongest points that make them a strong fit — draw from both CV analysis and LinkedIn
- Mention relevant domain experience, key skills, and any standout achievements visible on LinkedIn (publications, certifications, recommendations, awards if mentioned)
- If LinkedIn shows skill endorsements from peers, mention this as a credibility signal
- If call notes are provided, naturally weave in key insights — communication quality, salary expectation, joining timeline, relocation confirmation, motivation
- End with a strong, confident recommendation sentence stating why you are putting this candidate forward
- Tone: professional, confident, specific — avoid generic phrases like 'strong communication skills' unless supported by evidence
- Write in third person
- Do NOT mention red flags, concerns, or anything negative — this is a positive endorsement document
- Before writing, determine the candidate's gender from their name and any gender indicators in the CV. Use correct pronouns throughout — he/him for male candidates, she/her for female candidates. If gender cannot be determined, use their name instead of a pronoun. Never assume — always infer from the name or CV.
- Return only the remark text. No headings, no bullet points, plain paragraph."""

        response = client.chat.completions.create(
            model="llama-3.3-70b-versatile",
            messages=[
                {"role": "system", "content": system_msg},
                {"role": "user", "content": user_msg},
            ],
            temperature=0.4,
            max_tokens=600,
        )
        remark = response.choices[0].message.content.strip()
        return {"remark": remark}
    finally:
        db.close()


@app.put("/api/candidates/{candidate_id}/remark")
async def save_remark(candidate_id: int, req: RemarkSave):
    db = SessionLocal()
    try:
        candidate = db.query(DBCandidate).filter(DBCandidate.id == candidate_id).first()
        if not candidate:
            raise HTTPException(status_code=404, detail="Candidate not found")
        candidate.remark = req.remark
        candidate.linkedin_url = req.linkedin_url
        candidate.linkedin_text = req.linkedin_text
        db.commit()
        return {"ok": True}
    finally:
        db.close()


@app.patch("/api/candidates/{candidate_id}/status")
async def update_candidate_status(candidate_id: int, body: CandidateStatusUpdate):
    valid = {"pending", "shortlisted", "on_hold", "rejected"}
    if body.status not in valid:
        raise HTTPException(status_code=400, detail=f"Invalid status. Must be one of: {', '.join(sorted(valid))}")
    db = SessionLocal()
    try:
        candidate = db.query(DBCandidate).filter(DBCandidate.id == candidate_id).first()
        if not candidate:
            raise HTTPException(status_code=404, detail="Candidate not found")
        candidate.status = body.status
        db.commit()
        return {"success": True}
    finally:
        db.close()


VALID_PIPELINE_STAGES = {"screened", "submitted", "interview_scheduled", "offer_made", "joined", "dropped"}

@app.patch("/api/candidates/{candidate_id}/pipeline")
async def update_candidate_pipeline(candidate_id: int, body: PipelineStageUpdate):
    if body.stage not in VALID_PIPELINE_STAGES:
        raise HTTPException(status_code=400, detail=f"Invalid stage. Must be one of: {', '.join(sorted(VALID_PIPELINE_STAGES))}")
    db = SessionLocal()
    try:
        candidate = db.query(DBCandidate).filter(DBCandidate.id == candidate_id).first()
        if not candidate:
            raise HTTPException(status_code=404, detail="Candidate not found")
        candidate.pipeline_stage = body.stage
        db.commit()
        return {"success": True}
    finally:
        db.close()


@app.post("/api/candidates/{candidate_id}/parse-linkedin-pdf")
async def parse_linkedin_pdf(candidate_id: int, linkedin_pdf: UploadFile = File(...)):
    data = await linkedin_pdf.read()
    try:
        text = extract_pdf_text(data)
        if not text.strip():
            raise ValueError("empty")
        return {"linkedin_text": text}
    except Exception:
        raise HTTPException(status_code=400, detail="Could not extract text from PDF")



# ─── Admin Endpoints ──────────────────────────────────────────────────────────

@app.get("/api/admin/stats")
async def admin_stats():
    db = SessionLocal()
    try:
        return {
            "total_recruiters": db.query(Recruiter).count(),
            "total_projects": db.query(DBProject).count(),
            "total_candidates_screened": db.query(DBCandidate).count(),
            "total_shortlisted": db.query(DBCandidate).filter(DBCandidate.status == "shortlisted").count(),
            "total_on_hold": db.query(DBCandidate).filter(DBCandidate.status == "on_hold").count(),
            "total_rejected": db.query(DBCandidate).filter(DBCandidate.status == "rejected").count(),
            "total_pending": db.query(DBCandidate).filter(
                (DBCandidate.status == "pending") | (DBCandidate.status == None)
            ).count(),
            "total_submitted": db.query(DBCandidate).filter(DBCandidate.pipeline_stage == "submitted").count(),
            "total_interviewing": db.query(DBCandidate).filter(DBCandidate.pipeline_stage == "interview_scheduled").count(),
            "total_offers": db.query(DBCandidate).filter(DBCandidate.pipeline_stage == "offer_made").count(),
            "total_joined": db.query(DBCandidate).filter(DBCandidate.pipeline_stage == "joined").count(),
            "total_dropped": db.query(DBCandidate).filter(DBCandidate.pipeline_stage == "dropped").count(),
        }
    finally:
        db.close()


@app.get("/api/admin/recruiters")
async def admin_recruiters():
    db = SessionLocal()
    try:
        recruiters = db.query(Recruiter).all()
        result = []
        for r in recruiters:
            projects = db.query(DBProject).filter(DBProject.created_by_email == r.email).all()
            project_ids = [p.id for p in projects]
            if project_ids:
                candidates = db.query(DBCandidate).filter(DBCandidate.project_id.in_(project_ids)).all()
                candidates_screened = len(candidates)
                last_active = max((c.uploaded_at for c in candidates if c.uploaded_at), default=None)
                shortlisted = sum(1 for c in candidates if c.status == "shortlisted")
                on_hold     = sum(1 for c in candidates if c.status == "on_hold")
                rejected    = sum(1 for c in candidates if c.status == "rejected")
            else:
                candidates_screened = 0
                last_active = None
                shortlisted = on_hold = rejected = 0
            result.append({
                "id": r.id,
                "name": r.name,
                "email": r.email,
                "is_admin": r.is_admin or False,
                "created_at": r.created_at.isoformat() if r.created_at else None,
                "projects_count": len(projects),
                "candidates_screened": candidates_screened,
                "shortlisted": shortlisted,
                "on_hold": on_hold,
                "rejected": rejected,
                "last_active": last_active.isoformat() if last_active else None,
            })
        result.sort(key=lambda x: x["candidates_screened"], reverse=True)
        return result
    finally:
        db.close()


@app.get("/api/admin/projects")
async def admin_projects():
    db = SessionLocal()
    try:
        projects = db.query(DBProject).order_by(DBProject.created_at.desc()).all()
        result = []
        for p in projects:
            rec = db.query(Recruiter).filter(Recruiter.email == p.created_by_email).first()
            candidates_count = db.query(DBCandidate).filter(DBCandidate.project_id == p.id).count()
            result.append({
                "id": p.id,
                "name": p.name,
                "client_name": p.client_name or "",
                "created_by_email": p.created_by_email,
                "recruiter_name": rec.name if rec else p.created_by_email,
                "created_at": p.created_at.isoformat() if p.created_at else None,
                "candidates_count": candidates_count,
            })
        return result
    finally:
        db.close()


@app.post("/api/admin/deduplicate-candidates")
async def deduplicate_candidates_endpoint():
    db = SessionLocal()
    try:
        deleted = _deduplicate_candidates(db)
        return {"deleted": deleted}
    finally:
        db.close()


# ─── Invitation Email Helper ──────────────────────────────────────────────────

async def _send_invitation_email(to_email: str, to_name: str, invited_by_name: str, token: str) -> None:
    gmail_user = os.getenv("GMAIL_USER", "")
    gmail_pass = os.getenv("GMAIL_APP_PASSWORD", "")
    app_url = os.getenv("APP_URL", "http://localhost:3000")
    activation_link = f"{app_url}?token={token}"
    if not gmail_user or not gmail_pass:
        print(f"[DEV] Invitation for {to_email}: {activation_link}  (set GMAIL_USER + GMAIL_APP_PASSWORD to send real emails)")
        return
    msg = EmailMessage()
    msg["From"] = gmail_user
    msg["To"] = to_email
    msg["Subject"] = "You're invited to ResumeAI — Agile Technology Solutions"
    msg.set_content(
        f"Hi {to_name},\n\n"
        "You have been invited to join ResumeAI, the AI-powered recruitment platform at Agile Technology Solutions.\n\n"
        "Click the link below to activate your account:\n"
        f"{activation_link}\n\n"
        "This link expires in 7 days.\n\n"
        f"— {invited_by_name}\n"
        "Agile Technology Solutions"
    )
    await aiosmtplib.send(
        msg,
        hostname="smtp.gmail.com",
        port=587,
        start_tls=True,
        username=gmail_user,
        password=gmail_pass,
    )


# ─── Access Control Endpoints ─────────────────────────────────────────────────

@app.post("/api/admin/invite")
async def invite_user(body: InviteRequest):
    requester_email = body.requester_email.lower().strip()
    invite_email = body.email.lower().strip()

    db = SessionLocal()
    try:
        # Validate requester permissions
        if requester_email != ADMIN_EMAIL.lower().strip():
            requester = db.query(Recruiter).filter(Recruiter.email == requester_email).first()
            if not requester or requester.role not in ('admin', 'manager'):
                raise HTTPException(status_code=403, detail="Not authorised to invite users")
            if requester.role == 'manager' and body.role != 'recruiter':
                raise HTTPException(status_code=403, detail="Managers can only invite recruiters")

        # Get requester name for email
        req_rec = db.query(Recruiter).filter(Recruiter.email == requester_email).first()
        invited_by_name = req_rec.name if req_rec else requester_email

        # Upsert invitation
        token = secrets.token_urlsafe(32)
        existing_inv = db.query(Invitation).filter(Invitation.email == invite_email).first()
        if existing_inv:
            existing_inv.name = body.name or existing_inv.name
            existing_inv.role = body.role
            existing_inv.invited_by = requester_email
            existing_inv.team_manager_email = body.team_manager_email or None
            existing_inv.token = token
            existing_inv.status = 'pending'
            existing_inv.created_at = datetime.utcnow()
            existing_inv.accepted_at = None
        else:
            db.add(Invitation(
                email=invite_email,
                name=body.name or None,
                role=body.role,
                invited_by=requester_email,
                team_manager_email=body.team_manager_email or None,
                token=token,
            ))
        db.commit()
    finally:
        db.close()

    await _send_invitation_email(invite_email, body.name or invite_email, invited_by_name, token)
    return {"success": True}


@app.get("/api/admin/activate")
async def validate_invite_token(token: str):
    db = SessionLocal()
    try:
        inv = db.query(Invitation).filter(Invitation.token == token, Invitation.status == 'pending').first()
        if not inv:
            raise HTTPException(status_code=404, detail="This invitation link is invalid or has expired.")
        # Check 7-day expiry
        if inv.created_at and (datetime.utcnow() - inv.created_at).days >= 7:
            raise HTTPException(status_code=410, detail="This invitation link has expired. Please ask your admin for a new invitation.")
        return {"email": inv.email, "name": inv.name or "", "role": inv.role}
    finally:
        db.close()


@app.post("/api/admin/activate")
async def activate_invite(body: ActivateRequest):
    db = SessionLocal()
    try:
        inv = db.query(Invitation).filter(Invitation.token == body.token, Invitation.status == 'pending').first()
        if not inv:
            raise HTTPException(status_code=404, detail="This invitation link is invalid or has expired.")
        if inv.created_at and (datetime.utcnow() - inv.created_at).days >= 7:
            raise HTTPException(status_code=410, detail="This invitation link has expired. Please ask your admin for a new invitation.")

        # Create or update recruiter
        recruiter = db.query(Recruiter).filter(Recruiter.email == inv.email).first()
        if not recruiter:
            recruiter = Recruiter(
                name=inv.name or inv.email.split("@")[0],
                email=inv.email,
                role=inv.role,
                team_manager_email=inv.team_manager_email,
                status='active',
                invited_by=inv.invited_by,
                invited_at=datetime.utcnow(),
            )
            db.add(recruiter)
        else:
            recruiter.role = inv.role
            recruiter.team_manager_email = inv.team_manager_email
            recruiter.status = 'active'
            recruiter.invited_by = inv.invited_by
            if inv.name:
                recruiter.name = inv.name

        inv.status = 'accepted'
        inv.accepted_at = datetime.utcnow()
        db.commit()
        db.refresh(recruiter)

        return {"success": True, "recruiter": {
            "id": recruiter.id,
            "email": recruiter.email,
            "name": recruiter.name,
            "is_admin": recruiter.is_admin or False,
            "role": recruiter.role or 'recruiter',
        }}
    finally:
        db.close()


@app.get("/api/admin/invitations")
async def list_invitations():
    db = SessionLocal()
    try:
        invitations = db.query(Invitation).filter(Invitation.status == 'pending').order_by(Invitation.created_at.desc()).all()
        return [
            {
                "id": inv.id,
                "email": inv.email,
                "name": inv.name or "",
                "role": inv.role,
                "invited_by": inv.invited_by,
                "team_manager_email": inv.team_manager_email or "",
                "created_at": inv.created_at.isoformat() if inv.created_at else None,
            }
            for inv in invitations
        ]
    finally:
        db.close()


@app.delete("/api/admin/invitations/{invitation_id}")
async def cancel_invitation(invitation_id: int):
    db = SessionLocal()
    try:
        inv = db.query(Invitation).filter(Invitation.id == invitation_id).first()
        if not inv:
            raise HTTPException(status_code=404, detail="Invitation not found")
        db.delete(inv)
        db.commit()
        return {"success": True}
    finally:
        db.close()


@app.get("/api/admin/team")
async def admin_team():
    db = SessionLocal()
    try:
        recruiters = db.query(Recruiter).all()
        result = []
        for r in recruiters:
            projects = db.query(DBProject).filter(DBProject.created_by_email == r.email).all()
            project_ids = [p.id for p in projects]
            candidates_screened = shortlisted = 0
            last_active = None
            if project_ids:
                candidates = db.query(DBCandidate).filter(DBCandidate.project_id.in_(project_ids)).all()
                candidates_screened = len(candidates)
                shortlisted = sum(1 for c in candidates if c.status == 'shortlisted')
                last_active = max((c.uploaded_at for c in candidates if c.uploaded_at), default=None)
            effective_role = r.role or ('admin' if r.is_admin else 'recruiter')
            result.append({
                "id": r.id,
                "name": r.name,
                "email": r.email,
                "role": effective_role,
                "team_manager_email": r.team_manager_email or "",
                "status": r.status or 'active',
                "invited_by": r.invited_by or "",
                "is_admin": r.is_admin or False,
                "created_at": r.created_at.isoformat() if r.created_at else None,
                "projects_count": len(projects),
                "candidates_screened": candidates_screened,
                "shortlisted": shortlisted,
                "last_active": last_active.isoformat() if last_active else None,
            })
        result.sort(key=lambda x: x["candidates_screened"], reverse=True)
        return result
    finally:
        db.close()


@app.patch("/api/admin/recruiters/{recruiter_id}")
async def update_recruiter(recruiter_id: int, body: RecruiterUpdate):
    db = SessionLocal()
    try:
        recruiter = db.query(Recruiter).filter(Recruiter.id == recruiter_id).first()
        if not recruiter:
            raise HTTPException(status_code=404, detail="Recruiter not found")
        if body.status is not None:
            if body.status not in ('active', 'suspended'):
                raise HTTPException(status_code=400, detail="Invalid status — must be 'active' or 'suspended'")
            recruiter.status = body.status
        if body.role is not None:
            if body.role not in ('admin', 'manager', 'recruiter'):
                raise HTTPException(status_code=400, detail="Invalid role — must be 'admin', 'manager', or 'recruiter'")
            recruiter.role = body.role
            recruiter.is_admin = (body.role == 'admin')
        db.commit()
        return {"success": True}
    finally:
        db.close()


@app.get("/api/admin/weekly-activity")
async def weekly_activity():
    db = SessionLocal()
    try:
        result = []
        today = datetime.utcnow().date()
        day_names = ['Mon', 'Tue', 'Wed', 'Thu', 'Fri', 'Sat', 'Sun']
        for i in range(6, -1, -1):
            d = today - timedelta(days=i)
            start_dt = datetime(d.year, d.month, d.day)
            end_dt = start_dt + timedelta(days=1)
            count = db.query(DBCandidate).filter(
                DBCandidate.uploaded_at >= start_dt,
                DBCandidate.uploaded_at < end_dt,
            ).count()
            result.append({"date": day_names[d.weekday()], "count": count})
        return result
    finally:
        db.close()


@app.get("/api/admin/reports")
async def admin_reports():
    db = SessionLocal()
    try:
        # ── Recruiter summary ──────────────────────────────────────────────
        recruiters = db.query(Recruiter).all()
        recruiter_summary = []
        for r in recruiters:
            projects = db.query(DBProject).filter(DBProject.created_by_email == r.email).all()
            project_ids = [p.id for p in projects]
            candidates = db.query(DBCandidate).filter(DBCandidate.project_id.in_(project_ids)).all() if project_ids else []
            screened    = len(candidates)
            shortlisted = sum(1 for c in candidates if c.status == "shortlisted")
            on_hold     = sum(1 for c in candidates if c.status == "on_hold")
            rejected    = sum(1 for c in candidates if c.status == "rejected")
            submitted   = sum(1 for c in candidates if c.pipeline_stage == "submitted")
            interviewed = sum(1 for c in candidates if c.pipeline_stage == "interview_scheduled")
            offered     = sum(1 for c in candidates if c.pipeline_stage == "offer_made")
            joined      = sum(1 for c in candidates if c.pipeline_stage == "joined")
            dropped     = sum(1 for c in candidates if c.pipeline_stage == "dropped")
            scores      = [c.score for c in candidates if c.score is not None]
            avg_score   = round(sum(scores) / len(scores), 1) if scores else 0
            recruiter_summary.append({
                "recruiter_name": r.name,
                "recruiter_email": r.email,
                "total_screened": screened,
                "shortlisted": shortlisted,
                "on_hold": on_hold,
                "rejected": rejected,
                "submitted": submitted,
                "interviewed": interviewed,
                "offered": offered,
                "joined": joined,
                "dropped": dropped,
                "shortlist_rate": round(shortlisted / screened * 100, 1) if screened else 0,
                "conversion_rate": round(joined / screened * 100, 1) if screened else 0,
                "avg_score": avg_score,
                "projects_count": len(projects),
            })
        recruiter_summary.sort(key=lambda x: x["total_screened"], reverse=True)

        # ── Client summary ─────────────────────────────────────────────────
        all_projects = db.query(DBProject).all()
        client_map: dict = {}
        for p in all_projects:
            client = p.client_name or "Unknown"
            if client not in client_map:
                client_map[client] = {"client_name": client, "total_projects": 0, "total_screened": 0, "shortlisted": 0, "joined": 0}
            client_map[client]["total_projects"] += 1
            cands = db.query(DBCandidate).filter(DBCandidate.project_id == p.id).all()
            client_map[client]["total_screened"] += len(cands)
            client_map[client]["shortlisted"]    += sum(1 for c in cands if c.status == "shortlisted")
            client_map[client]["joined"]          += sum(1 for c in cands if c.pipeline_stage == "joined")
        client_summary = sorted(
            [
                {**d, "conversion_rate": round(d["joined"] / d["total_screened"] * 100, 1) if d["total_screened"] else 0}
                for d in client_map.values()
            ],
            key=lambda x: x["total_screened"],
            reverse=True,
        )

        # ── Pipeline summary ───────────────────────────────────────────────
        all_candidates = db.query(DBCandidate).all()
        pipeline_summary = {
            "screened":    len(all_candidates),
            "shortlisted": sum(1 for c in all_candidates if c.status == "shortlisted"),
            "submitted":   sum(1 for c in all_candidates if c.pipeline_stage == "submitted"),
            "interviewed": sum(1 for c in all_candidates if c.pipeline_stage == "interview_scheduled"),
            "offered":     sum(1 for c in all_candidates if c.pipeline_stage == "offer_made"),
            "joined":      sum(1 for c in all_candidates if c.pipeline_stage == "joined"),
            "dropped":     sum(1 for c in all_candidates if c.pipeline_stage == "dropped"),
        }

        # ── Top performers ─────────────────────────────────────────────────
        top_performers = []
        active = [r for r in recruiter_summary if r["total_screened"] > 0]
        if active:
            m = max(active, key=lambda x: x["total_screened"])
            top_performers.append({"recruiter_name": m["recruiter_name"], "metric": "Most Active", "value": f"{m['total_screened']} screened"})
            m = max(active, key=lambda x: x["shortlist_rate"])
            top_performers.append({"recruiter_name": m["recruiter_name"], "metric": "Best Shortlister", "value": f"{m['shortlist_rate']}% shortlist rate"})
            m = max(active, key=lambda x: x["conversion_rate"])
            top_performers.append({"recruiter_name": m["recruiter_name"], "metric": "Best Converter", "value": f"{m['conversion_rate']}% conversion rate"})
            m = max(active, key=lambda x: x["joined"])
            top_performers.append({"recruiter_name": m["recruiter_name"], "metric": "Most Placements", "value": f"{m['joined']} joined"})

        return {
            "recruiter_summary": recruiter_summary,
            "client_summary": client_summary,
            "pipeline_summary": pipeline_summary,
            "top_performers": top_performers,
        }
    finally:
        db.close()


# ─── Client Learning Helper ───────────────────────────────────────────────────

_LEARNING_MAP = {
    "skills_mismatch":        "Candidates missing key skills are consistently rejected",
    "overqualified":          "Overqualified candidates are rejected — match seniority carefully",
    "company_tier_mismatch":  "Company tier is important for this client",
    "budget":                 "Budget constraints are common — verify CTC fit before submitting",
    "stability_concern":      "Client values stability — job hoppers are rejected",
    "domain_mismatch":        "Domain experience is non-negotiable for this client",
}

def _update_client_learning(client_name: str, reason: str, db) -> None:
    if not client_name or client_name.strip().lower() in ("", "unknown"):
        return
    if reason not in _LEARNING_MAP:
        return
    existing = db.query(ClientLearning).filter(
        ClientLearning.client_name == client_name,
        ClientLearning.learning_type == reason,
    ).first()
    if existing:
        existing.confidence += 1
        existing.updated_at = datetime.utcnow()
    else:
        db.add(ClientLearning(
            client_name=client_name,
            learning_type=reason,
            learning_text=_LEARNING_MAP[reason],
            confidence=1,
        ))
    db.commit()


# ─── Feedback & Learning Endpoints ────────────────────────────────────────────

@app.post("/api/candidates/{candidate_id}/feedback")
async def submit_candidate_feedback(candidate_id: int, body: FeedbackRequest):
    db = SessionLocal()
    try:
        db.add(CandidateFeedback(
            candidate_id=candidate_id,
            project_id=body.project_id,
            recruiter_email=body.recruiter_email,
            action=body.action,
            reason=body.reason,
            reason_detail=body.reason_detail or "",
        ))
        db.commit()

        # Trigger client learning for rejection/drop actions
        if body.action in ("rejected", "dropped"):
            project = db.query(DBProject).filter(DBProject.id == body.project_id).first()
            client_name = project.client_name if project else None
            if client_name:
                _update_client_learning(client_name, body.reason, db)

        return {"success": True}
    finally:
        db.close()


@app.get("/api/client-learning/{client_name}")
async def get_client_learning(client_name: str):
    db = SessionLocal()
    try:
        rules = db.query(ClientLearning).filter(
            ClientLearning.client_name == client_name,
        ).order_by(ClientLearning.confidence.desc()).all()
        return [
            {
                "learning_type": r.learning_type,
                "learning_text": r.learning_text,
                "confidence": r.confidence,
            }
            for r in rules
        ]
    finally:
        db.close()


@app.get("/api/admin/client-learning")
async def admin_client_learning():
    db = SessionLocal()
    try:
        rules = db.query(ClientLearning).order_by(
            ClientLearning.client_name, ClientLearning.confidence.desc()
        ).all()
        grouped: dict = {}
        for r in rules:
            grouped.setdefault(r.client_name, []).append({
                "learning_type": r.learning_type,
                "learning_text": r.learning_text,
                "confidence": r.confidence,
            })
        return [{"client_name": k, "rules": v} for k, v in grouped.items()]
    finally:
        db.close()


# ─── Feature Usage Endpoints ──────────────────────────────────────────────────

@app.post("/api/feature-usage")
async def track_feature_usage(body: FeatureUsageRequest):
    db = SessionLocal()
    try:
        db.add(FeatureUsage(
            recruiter_email=body.recruiter_email,
            feature_name=body.feature_name,
        ))
        db.commit()
        return {"success": True}
    finally:
        db.close()


@app.get("/api/admin/feature-usage")
async def admin_feature_usage():
    db = SessionLocal()
    try:
        from sqlalchemy import func
        rows = (
            db.query(FeatureUsage.feature_name, func.count(FeatureUsage.id).label("count"))
            .group_by(FeatureUsage.feature_name)
            .order_by(func.count(FeatureUsage.id).desc())
            .all()
        )
        all_features = [{"feature": r.feature_name, "count": r.count} for r in rows]
        most_used  = all_features[:10]
        least_used = list(reversed(all_features))[:5]

        by_recruiter_rows = (
            db.query(
                FeatureUsage.recruiter_email,
                func.count(FeatureUsage.id).label("features_used"),
            )
            .group_by(FeatureUsage.recruiter_email)
            .order_by(func.count(FeatureUsage.id).desc())
            .all()
        )
        by_recruiter = []
        for row in by_recruiter_rows:
            rec = db.query(Recruiter).filter(Recruiter.email == row.recruiter_email).first()
            by_recruiter.append({
                "email": row.recruiter_email,
                "name": rec.name if rec else row.recruiter_email,
                "features_used": row.features_used,
            })

        return {"most_used": most_used, "least_used": least_used, "by_recruiter": by_recruiter}
    finally:
        db.close()


# ─── Survey Endpoints ─────────────────────────────────────────────────────────

@app.get("/api/survey/check")
async def survey_check(email: str):
    month = datetime.utcnow().strftime("%Y-%m")
    db = SessionLocal()
    try:
        existing = db.query(MonthlySurvey).filter(
            MonthlySurvey.recruiter_email == email,
            MonthlySurvey.month == month,
        ).first()
        return {"completed": existing is not None, "month": month}
    finally:
        db.close()


@app.post("/api/survey/submit")
async def survey_submit(body: SurveyRequest):
    month = datetime.utcnow().strftime("%Y-%m")
    db = SessionLocal()
    try:
        existing = db.query(MonthlySurvey).filter(
            MonthlySurvey.recruiter_email == body.recruiter_email,
            MonthlySurvey.month == month,
        ).first()
        if existing:
            existing.q1_rating = body.q1_rating
            existing.q2_rating = body.q2_rating
            existing.q3_text = body.q3_text
            existing.completed_at = datetime.utcnow()
        else:
            db.add(MonthlySurvey(
                recruiter_email=body.recruiter_email,
                month=month,
                q1_rating=body.q1_rating,
                q2_rating=body.q2_rating,
                q3_text=body.q3_text,
            ))
        db.commit()
        return {"success": True}
    finally:
        db.close()


# ─── Auth Endpoints ───────────────────────────────────────────────────────────

async def _send_otp_email(to_email: str, code: str) -> None:
    gmail_user = os.getenv("GMAIL_USER", "")
    gmail_pass = os.getenv("GMAIL_APP_PASSWORD", "")
    if not gmail_user or not gmail_pass:
        print(f"[DEV] OTP for {to_email}: {code}  (set GMAIL_USER + GMAIL_APP_PASSWORD to send real emails)")
        return
    msg = EmailMessage()
    msg["From"] = gmail_user
    msg["To"] = to_email
    msg["Subject"] = "Your ResumeAI login code"
    msg.set_content(
        f"Your login code is: {code}\n\n"
        "This code expires in 10 minutes.\n"
        "Do not share this code with anyone."
    )
    await aiosmtplib.send(
        msg,
        hostname="smtp.gmail.com",
        port=587,
        start_tls=True,
        username=gmail_user,
        password=gmail_pass,
    )


@app.post("/api/auth/request-otp")
async def request_otp(body: OTPRequest):
    body.email = body.email.lower().strip()
    db = SessionLocal()
    try:
        # Access control — only allowed users can log in
        if body.email != ADMIN_EMAIL.lower().strip():
            recruiter_record = db.query(Recruiter).filter(Recruiter.email == body.email).first()
            if not recruiter_record or recruiter_record.status != 'active':
                # Check if they have an accepted invitation (fallback)
                accepted_inv = db.query(Invitation).filter(
                    Invitation.email == body.email,
                    Invitation.status == 'accepted',
                ).first()
                if not accepted_inv:
                    raise HTTPException(
                        status_code=403,
                        detail={
                            "error": "access_denied",
                            "message": "You are not authorised to access ResumeAI. Please ask your admin to invite you.",
                        },
                    )

        # Upsert recruiter
        recruiter = db.query(Recruiter).filter(Recruiter.email == body.email).first()
        if not recruiter:
            recruiter = Recruiter(name=body.name or body.email.split("@")[0], email=body.email)
            db.add(recruiter)
        elif body.name:
            recruiter.name = body.name
        if body.email == ADMIN_EMAIL.lower().strip():
            recruiter.is_admin = True
            recruiter.role = 'admin'
        db.commit()

        # Generate and save OTP
        code = str(random.randint(100000, 999999))
        now = datetime.utcnow()
        otp = OTPCode(email=body.email, code=code, created_at=now, expires_at=now + timedelta(minutes=10))
        db.add(otp)
        db.commit()
    finally:
        db.close()

    await _send_otp_email(body.email, code)
    return {"message": "OTP sent"}


@app.post("/api/auth/verify-otp")
async def verify_otp(body: OTPVerify):
    body.email = body.email.lower().strip()
    db = SessionLocal()
    try:
        now = datetime.utcnow()
        otp = (
            db.query(OTPCode)
            .filter(
                OTPCode.email == body.email,
                OTPCode.code == body.code,
                OTPCode.used == False,
                OTPCode.expires_at > now,
            )
            .order_by(OTPCode.created_at.desc())
            .first()
        )
        if not otp:
            return {"success": False, "message": "Invalid or expired code"}

        otp.used = True
        db.commit()

        recruiter = db.query(Recruiter).filter(Recruiter.email == body.email).first()
        if not recruiter:
            return {"success": False, "message": "Recruiter not found"}

        effective_role = recruiter.role or ('admin' if recruiter.is_admin else 'recruiter')
        return {"success": True, "recruiter": {
            "id": recruiter.id,
            "email": recruiter.email,
            "name": recruiter.name,
            "is_admin": recruiter.is_admin or False,
            "role": effective_role,
        }}
    finally:
        db.close()


@app.post("/api/auth/logout")
async def logout():
    return {"success": True}


# ─── Projects Endpoints ───────────────────────────────────────────────────────

@app.get("/api/projects")
async def list_projects():
    db = SessionLocal()
    try:
        projects = db.query(DBProject).order_by(DBProject.created_at.desc()).all()
        result = []
        for p in projects:
            count = db.query(DBCandidate).filter(DBCandidate.project_id == p.id).count()
            result.append({
                "id": p.id,
                "name": p.name,
                "client_name": p.client_name,
                "created_by_email": p.created_by_email,
                "created_at": p.created_at.isoformat(),
                "candidate_count": count,
                "has_jd": bool(p.jd_text),
            })
        return result
    finally:
        db.close()


@app.post("/api/projects")
async def create_project(body: ProjectCreate):
    db = SessionLocal()
    try:
        project = DBProject(
            name=body.name,
            client_name=body.client_name or None,
            created_by_email=body.created_by_email,
        )
        db.add(project)
        db.commit()
        db.refresh(project)
        return {
            "id": project.id,
            "name": project.name,
            "client_name": project.client_name,
            "created_by_email": project.created_by_email,
            "created_at": project.created_at.isoformat(),
            "candidate_count": 0,
        }
    finally:
        db.close()


@app.get("/api/projects/{project_id}")
async def get_project(project_id: int):
    db = SessionLocal()
    try:
        project = db.query(DBProject).filter(DBProject.id == project_id).first()
        if not project:
            raise HTTPException(status_code=404, detail="Project not found")
        candidates = (
            db.query(DBCandidate)
            .filter(DBCandidate.project_id == project_id)
            .order_by(DBCandidate.score.desc())
            .all()
        )
        return {
            "id": project.id,
            "name": project.name,
            "client_name": project.client_name,
            "jd_text": project.jd_text,
            "jd_context": project.jd_context,
            "required_location": project.required_location,
            "keywords": project.keywords,
            "benchmark_fingerprint": project.benchmark_fingerprint,
            "created_by_email": project.created_by_email,
            "created_at": project.created_at.isoformat(),
            "candidates": [
                {
                    "id": c.id,
                    "name": c.name,
                    "filename": c.filename,
                    "score": c.score,
                    "result_json": c.result_json,
                    "status": c.status or "pending",
                    "remark": c.remark,
                    "linkedin_url": c.linkedin_url,
                    "linkedin_text": c.linkedin_text,
                    "benchmark_match_percent": c.benchmark_match_percent,
                    "benchmark_similarities": c.benchmark_similarities,
                    "benchmark_differences": c.benchmark_differences,
                    "benchmark_verdict": c.benchmark_verdict,
                    "candidate_phone": c.candidate_phone,
                    "is_duplicate": c.is_duplicate or False,
                    "duplicate_projects": json.loads(c.duplicate_projects) if c.duplicate_projects else [],
                    "pipeline_stage": c.pipeline_stage or "screened",
                    "uploaded_at": c.uploaded_at.isoformat(),
                }
                for c in candidates
            ],
        }
    finally:
        db.close()


@app.put("/api/projects/{project_id}/jd-context")
async def update_jd_context(project_id: int, body: JDContextUpdate):
    db = SessionLocal()
    try:
        project = db.query(DBProject).filter(DBProject.id == project_id).first()
        if not project:
            raise HTTPException(status_code=404, detail="Project not found")
        if body.jd_text:
            project.jd_text = body.jd_text
        if body.jd_context:
            project.jd_context = body.jd_context
        if body.required_location:
            project.required_location = body.required_location
        if body.keywords:
            project.keywords = body.keywords
        db.commit()
        return {"success": True}
    finally:
        db.close()


@app.post("/api/projects/{project_id}/candidates")
async def save_candidate(project_id: int, body: CandidateSave):
    db = SessionLocal()
    try:
        project = db.query(DBProject).filter(DBProject.id == project_id).first()
        if not project:
            raise HTTPException(status_code=404, detail="Project not found")
        existing = db.query(DBCandidate).filter(
            DBCandidate.project_id == project_id,
            DBCandidate.filename == body.filename,
        ).first()
        if existing:
            existing.name = body.name
            existing.score = body.score
            existing.result_json = body.result_json
            existing.candidate_phone = body.candidate_phone
            saved_candidate = existing
        else:
            saved_candidate = DBCandidate(
                project_id=project_id,
                name=body.name,
                filename=body.filename,
                score=body.score,
                result_json=body.result_json,
                candidate_phone=body.candidate_phone,
            )
            db.add(saved_candidate)

        # Duplicate detection — find same candidate in other projects by name or phone
        dup_matches = []
        if body.name:
            name_q = db.query(DBCandidate).filter(
                DBCandidate.project_id != project_id,
                DBCandidate.name == body.name,
            )
            dup_matches.extend(name_q.all())
        if body.candidate_phone:
            phone_matches = db.query(DBCandidate).filter(
                DBCandidate.project_id != project_id,
                DBCandidate.candidate_phone == body.candidate_phone,
                DBCandidate.candidate_phone != None,
            ).all()
            existing_ids = {c.id for c in dup_matches}
            dup_matches.extend(c for c in phone_matches if c.id not in existing_ids)

        if dup_matches:
            dup_project_ids = list({c.project_id for c in dup_matches})
            dup_projects_data = []
            for pid in dup_project_ids:
                proj = db.query(DBProject).filter(DBProject.id == pid).first()
                if proj:
                    dup_projects_data.append({"project_id": pid, "project_name": proj.name})
            saved_candidate.is_duplicate = True
            saved_candidate.duplicate_projects = json.dumps(dup_projects_data)
            duplicate_info = dup_projects_data
        else:
            saved_candidate.is_duplicate = False
            saved_candidate.duplicate_projects = None
            duplicate_info = []

        db.commit()
        db.refresh(saved_candidate)
        return {"id": saved_candidate.id, "name": saved_candidate.name, "score": saved_candidate.score, "duplicate_info": duplicate_info}
    finally:
        db.close()


@app.delete("/api/projects/{project_id}")
async def delete_project(project_id: int):
    db = SessionLocal()
    try:
        project = db.query(DBProject).filter(DBProject.id == project_id).first()
        if not project:
            raise HTTPException(status_code=404, detail="Project not found")
        db.delete(project)
        db.commit()
        return {"success": True}
    finally:
        db.close()


# ─── Post-Screening Chat ──────────────────────────────────────────────────────

@app.post("/api/projects/{project_id}/chat")
async def chat_about_candidate(project_id: int, body: ChatRequest):
    db = SessionLocal()
    try:
        project = db.query(DBProject).filter(DBProject.id == project_id).first()
        if not project:
            raise HTTPException(status_code=404, detail="Project not found")

        candidate = db.query(DBCandidate).filter(
            DBCandidate.id == body.candidate_id,
            DBCandidate.project_id == project_id,
        ).first()
        if not candidate:
            raise HTTPException(status_code=404, detail="Candidate not found")

        result = json.loads(candidate.result_json or "{}")

        others = (
            db.query(DBCandidate)
            .filter(DBCandidate.project_id == project_id, DBCandidate.id != body.candidate_id)
            .order_by(DBCandidate.score.desc())
            .all()
        )

        try:
            kw_list = json.loads(project.keywords or "[]")
            keywords_str = ", ".join(kw_list) if kw_list else "Not specified"
        except Exception:
            keywords_str = project.keywords or "Not specified"

        comparison_lines = []
        for oc in others:
            or_ = json.loads(oc.result_json or "{}")
            comparison_lines.append(
                f"  - {oc.name or or_.get('candidate_name', 'Unknown')} | "
                f"Score: {oc.score}/100 | Recommendation: {or_.get('recommendation', 'N/A')}"
            )
        comparison_text = "\n".join(comparison_lines) if comparison_lines else "  No other candidates screened yet."

        cf = result.get("company_fit", {})
        jhf = result.get("job_hopping_flags", [])
        jh_text = ", ".join(jhf) if jhf else "None"
        gaps = result.get("career_gaps", [])
        gaps_text = (
            ", ".join(
                f"{g.get('from', '?')} to {g.get('to', '?')} ({g.get('gap_months', 0)}m)"
                for g in gaps
            )
            if gaps
            else "None"
        )
        edu = result.get("education_detail", result.get("education_level", "N/A"))

        system_prompt = f"""You are an expert recruitment advisor helping a recruiter understand a candidate's screening results.

You have full context about the job requirement, this candidate's CV analysis, and all other candidates screened for this role.

Answer the recruiter's questions clearly and concisely. Be direct and specific — use actual data from the analysis (scores, skill names, experience numbers, company names) in your answers.

Rules:
- Always refer to the candidate by first name
- If asked to compare, reference specific scores and differences
- If asked about red flags, be honest but constructive
- If asked for a recommendation, give a clear yes/no with reasoning
- Keep answers to 3-5 lines maximum unless the question needs more detail
- Use simple language — the recruiter may not be technical

JOB DETAILS:
Required Location: {project.required_location or 'Not specified'}
Key Skills: {keywords_str}
JD Summary: {(project.jd_text or '')[:500]}

CANDIDATE BEING DISCUSSED:
Name: {result.get('candidate_name', candidate.name)}
Score: {candidate.score}/100
Recommendation: {result.get('recommendation', 'N/A')}
Recommendation Reason: {result.get('recommendation_reason', 'N/A')}
Total Experience: {result.get('total_experience_display', 'N/A')}
Relevant Experience: {result.get('relevant_experience_display', 'N/A')}
Location: {result.get('candidate_location', 'N/A')}
Location Flag: {result.get('location_flag', 'N/A')} — {result.get('location_reason', '')}
Education: {edu}
Company Fit: {cf.get('company_fit_verdict', 'N/A')} ({cf.get('company_fit_score', 'N/A')}/10) — {cf.get('company_fit_explanation', '')}
Job Hopping: {jh_text}
Career Gaps: {gaps_text}
Matched Skills: {', '.join(result.get('matched_skills', [])) or 'None'}
Missing Skills: {', '.join(result.get('missing_skills', [])) or 'None'}
Current Company: {cf.get('candidate_current_company', 'N/A')}
AI Summary: {result.get('summary', 'N/A')}
Strengths: {', '.join(result.get('strengths', [])) or 'None'}
Concerns: {', '.join(result.get('concerns', [])) or 'None'}

OTHER CANDIDATES IN THIS PROJECT (for comparison):
{comparison_text}

Answer the recruiter's question based on this data."""

        messages = [{"role": "system", "content": system_prompt}]
        for msg in body.chat_history:
            messages.append({"role": msg.role, "content": msg.content})
        messages.append({"role": "user", "content": body.message})

        response = client.chat.completions.create(
            model="llama-3.3-70b-versatile",
            max_tokens=500,
            temperature=0.7,
            messages=messages,
        )

        return {
            "response": response.choices[0].message.content.strip(),
            "candidate_name": result.get("candidate_name", candidate.name),
        }
    finally:
        db.close()


# ─── Daily Content ────────────────────────────────────────────────────────────

DAILY_FALLBACK = {
    "quote": "Every great hire starts with a great conversation.",
    "quote_author": "ResumeAI Insight",
    "quote_category": "motivational",
    "trivia": "Did you know? The average recruiter spends 6 seconds scanning a resume before deciding to read further.",
}

@app.get("/api/daily-content")
async def get_daily_content():
    today = datetime.utcnow().strftime("%Y-%m-%d")
    db = SessionLocal()
    try:
        prompt = (
            f"Generate fresh daily content for a recruitment platform used by Indian recruiters. "
            f"Return JSON with exactly these fields: "
            f"quote (inspiring quote about hiring/talent/careers, max 2 sentences), "
            f"quote_author (person or 'ResumeAI Insight' if original), "
            f"quote_category (one of: motivational, trivia, fun), "
            f"trivia (HR fact for Indian market, start with 'Did you know?', max 2 sentences). "
            f"Today is {today}. Return only valid JSON, no markdown."
        )
        try:
            resp = client.chat.completions.create(
                model="llama-3.3-70b-versatile",
                max_tokens=300,
                temperature=0.9,
                messages=[{"role": "user", "content": prompt}],
            )
            raw = resp.choices[0].message.content.strip()
            data = json.loads(raw)
            existing = db.query(DailyContent).filter(DailyContent.date == today).first()
            if existing:
                existing.quote = data["quote"]
                existing.quote_author = data["quote_author"]
                existing.quote_category = data["quote_category"]
                existing.trivia = data["trivia"]
            else:
                db.add(DailyContent(
                    date=today,
                    quote=data["quote"],
                    quote_author=data["quote_author"],
                    quote_category=data["quote_category"],
                    trivia=data["trivia"],
                ))
            db.commit()
            return data
        except Exception:
            return DAILY_FALLBACK
    finally:
        db.close()


# ─── Games ────────────────────────────────────────────────────────────────────

async def _get_or_generate_game(game_type: str, today_str: str, db) -> dict:
    existing = db.query(DailyGame).filter(
        DailyGame.game_type == game_type,
        DailyGame.game_date == today_str,
    ).first()
    if existing:
        return json.loads(existing.content)

    if game_type == 'sliding_puzzle':
        content = {"size": 3}
    elif game_type == 'mcq':
        try:
            resp = client.chat.completions.create(
                model="llama-3.3-70b-versatile",
                max_tokens=400,
                temperature=0.8,
                messages=[{
                    "role": "user",
                    "content": (
                        "Generate a recruitment trivia question for Indian recruiters. "
                        "Topics: hiring best practices, HR law India, interview techniques, ATS, employer branding. "
                        "Return ONLY valid JSON, no markdown:\n"
                        '{"question":"...","options":["A. ...","B. ...","C. ...","D. ..."],"correct":0,"explanation":"..."}\n'
                        "correct is the 0-based index of the correct option."
                    ),
                }],
            )
            raw = resp.choices[0].message.content.strip()
            raw = re.sub(r"^```(?:json)?\s*", "", raw, flags=re.MULTILINE)
            raw = re.sub(r"\s*```$", "", raw, flags=re.MULTILINE)
            content = json.loads(raw)
        except Exception:
            content = {
                "question": "What is the recommended notice period for most mid-level roles in India?",
                "options": ["A. 15 days", "B. 30 days", "C. 60 days", "D. 90 days"],
                "correct": 1,
                "explanation": "30 days is the standard notice period for most mid-level roles, though senior roles often require 60–90 days.",
            }
    elif game_type == 'connections':
        try:
            resp = client.chat.completions.create(
                model="llama-3.3-70b-versatile",
                max_tokens=600,
                temperature=0.8,
                messages=[{
                    "role": "user",
                    "content": (
                        "Create a Word Connections puzzle about recruitment and HR for Indian recruiters. "
                        "4 groups of 4 words each, all related to hiring. "
                        "Return ONLY valid JSON, no markdown:\n"
                        '{"groups":[{"category":"...","words":["","","",""],"difficulty":1},{"category":"...","words":["","","",""],"difficulty":2},{"category":"...","words":["","","",""],"difficulty":3},{"category":"...","words":["","","",""],"difficulty":4}],"words":["all","16","words","shuffled","randomly","here","fill","rest"]}\n'
                        "difficulty 1=easy(yellow), 2=medium(green), 3=hard(blue), 4=tricky(purple). "
                        "words array must contain all 16 words shuffled randomly."
                    ),
                }],
            )
            raw = resp.choices[0].message.content.strip()
            raw = re.sub(r"^```(?:json)?\s*", "", raw, flags=re.MULTILINE)
            raw = re.sub(r"\s*```$", "", raw, flags=re.MULTILINE)
            content = json.loads(raw)
            all_words = [w for g in content["groups"] for w in g["words"]]
            random.shuffle(all_words)
            content["words"] = all_words
        except Exception:
            content = {
                "groups": [
                    {"category": "Interview stages", "words": ["Screening", "Technical", "HR Round", "Offer"], "difficulty": 1},
                    {"category": "Job portals India", "words": ["Naukri", "LinkedIn", "Indeed", "Shine"], "difficulty": 2},
                    {"category": "Resume sections", "words": ["Summary", "Skills", "Education", "Projects"], "difficulty": 3},
                    {"category": "Hiring red flags", "words": ["Job hopping", "Gap year", "Overqualified", "Relocation"], "difficulty": 4},
                ],
                "words": ["Screening", "Technical", "HR Round", "Offer", "Naukri", "LinkedIn", "Indeed", "Shine",
                          "Summary", "Skills", "Education", "Projects", "Job hopping", "Gap year", "Overqualified", "Relocation"],
            }
            random.shuffle(content["words"])
    elif game_type == 'red_flag':
        try:
            resp = client.chat.completions.create(
                model="llama-3.3-70b-versatile",
                max_tokens=600,
                temperature=0.8,
                messages=[{
                    "role": "user",
                    "content": (
                        "Create a Resume Red Flag challenge for Indian recruiters. "
                        "Write a realistic 3-4 line resume snippet for an Indian job applicant that contains ONE subtle red flag "
                        "(e.g. unexplained gap, job hopping, inconsistent dates, inflated title, missing education). "
                        "Provide 4 options where only one correctly identifies the red flag. "
                        "Return ONLY valid JSON, no markdown:\n"
                        '{"snippet":"resume text here...","options":["A. ...","B. ...","C. ...","D. ..."],"correct":0,"explanation":"why this is the red flag..."}\n'
                        "correct is the 0-based index. Other options should be plausible but incorrect."
                    ),
                }],
            )
            raw = resp.choices[0].message.content.strip()
            raw = re.sub(r"^```(?:json)?\s*", "", raw, flags=re.MULTILINE)
            raw = re.sub(r"\s*```$", "", raw, flags=re.MULTILINE)
            content = json.loads(raw)
        except Exception:
            content = {
                "snippet": "Priya Sharma | Sales Manager at TechCorp (Jan 2022 – Present)\nPreviously: Sales Executive at AlphaSoft (Mar 2021 – Dec 2021), Sales Intern at BetaInc (Jun 2020 – Feb 2021)\nBE Computer Science, Pune University, 2020 | Notice Period: Immediate",
                "options": ["A. No LinkedIn profile mentioned", "B. Three roles in under 3 years with short tenures", "C. Immediate joiner is suspicious", "D. Internship listed on resume"],
                "correct": 1,
                "explanation": "Three positions in under 3 years, with the earliest two lasting under 12 months each, is a classic job-hopping pattern and a key screening red flag.",
            }
    elif game_type == 'jd_match':
        try:
            resp = client.chat.completions.create(
                model="llama-3.3-70b-versatile",
                max_tokens=700,
                temperature=0.8,
                messages=[{
                    "role": "user",
                    "content": (
                        "Create a JD Match puzzle for Indian recruiters. "
                        "Generate: (1) a short realistic Job Description for an Indian company (4-6 lines), "
                        "(2) exactly 3 candidate profiles with Indian names (2-3 lines each — include role, experience, key skills), "
                        "(3) the 0-based index of the best-matching candidate, "
                        "(4) a 2-sentence explanation of why they are the best match. "
                        "Return ONLY valid JSON, no markdown:\n"
                        '{"jd":"...","candidates":[{"name":"...","profile":"..."},{"name":"...","profile":"..."},{"name":"...","profile":"..."}],"correct":1,"ai_reason":"..."}\n'
                        "Make the correct answer non-obvious — not always index 0."
                    ),
                }],
            )
            raw = resp.choices[0].message.content.strip()
            raw = re.sub(r"^```(?:json)?\s*", "", raw, flags=re.MULTILINE)
            raw = re.sub(r"\s*```$", "", raw, flags=re.MULTILINE)
            content = json.loads(raw)
        except Exception:
            content = {
                "jd": "We are hiring a Senior HR Recruiter for a mid-size IT company in Pune. The role requires 3-5 years of full-cycle recruitment experience, strong knowledge of ATS tools, and experience hiring for technical roles. Immediate joiners preferred.",
                "candidates": [
                    {"name": "Amit Verma", "profile": "5 yrs experience in FMCG recruitment. Strong sourcing skills via Naukri. No tech hiring experience. 30-day notice period."},
                    {"name": "Sneha Patil", "profile": "4 yrs IT recruitment at a Pune-based product company. Proficient in Greenhouse ATS. Hired 50+ engineers in last 2 yrs. Immediate joiner."},
                    {"name": "Ravi Kumar", "profile": "2 yrs general HR, currently in payroll operations. No recruitment experience. MBA in HR from Symbiosis."},
                ],
                "correct": 1,
                "ai_reason": "Sneha has 4 years of direct IT recruitment experience, ATS proficiency, and a proven track record of technical hiring — exactly what this role needs. Her immediate availability is an additional advantage.",
            }
    else:
        content = {}

    db.add(DailyGame(game_type=game_type, game_date=today_str, content=json.dumps(content)))
    db.commit()
    return content


@app.get("/api/games/daily")
async def get_daily_games(email: str):
    today_str = date.today().isoformat()
    db = SessionLocal()
    try:
        result = {}
        for game_type in ['mcq', 'connections', 'red_flag', 'jd_match', 'sliding_puzzle']:
            content = await _get_or_generate_game(game_type, today_str, db)
            played = db.query(GameResult).filter(
                GameResult.recruiter_email == email,
                GameResult.game_type == game_type,
                GameResult.game_date == today_str,
            ).first()
            result[game_type] = {
                "content": content,
                "played": bool(played),
                "score": played.score if played else 0,
            }
        streak = db.query(Streak).filter(Streak.recruiter_email == email).first()
        result["streak"] = {
            "current": streak.current_streak if streak else 0,
            "longest": streak.longest_streak if streak else 0,
            "total_games": streak.total_games if streak else 0,
        }
        return result
    finally:
        db.close()


@app.post("/api/games/result")
async def save_game_result(body: GameResultRequest):
    today_str = date.today().isoformat()
    yesterday_str = (date.today() - timedelta(days=1)).isoformat()
    db = SessionLocal()
    try:
        # Count games already played today BEFORE inserting this one
        games_played_today = db.query(GameResult).filter(
            GameResult.recruiter_email == body.email,
            GameResult.game_date == today_str,
        ).count()

        # Upsert game result (prevent duplicate plays)
        existing = db.query(GameResult).filter(
            GameResult.recruiter_email == body.email,
            GameResult.game_type == body.game_type,
            GameResult.game_date == today_str,
        ).first()
        if existing:
            existing.score = max(existing.score, body.score)
            existing.completed = body.completed or existing.completed
            if body.time_taken > 0:
                existing.time_taken = body.time_taken
            is_new_game = False
        else:
            db.add(GameResult(
                recruiter_email=body.email,
                game_type=body.game_type,
                game_date=today_str,
                score=body.score,
                completed=body.completed,
                time_taken=body.time_taken,
            ))
            is_new_game = True

        # Update streak — only counts when recruiter plays >= 2 games in a day
        streak = db.query(Streak).filter(Streak.recruiter_email == body.email).first()
        if not streak:
            streak = Streak(recruiter_email=body.email, current_streak=0, longest_streak=0, total_games=0)
            db.add(streak)
            db.flush()

        if is_new_game:
            streak.total_games += 1
            # Streak increments on the 2nd distinct game of the day
            should_advance = (games_played_today + 1) >= 2
            if should_advance and streak.last_played_date != today_str:
                if streak.last_played_date == yesterday_str:
                    streak.current_streak += 1
                else:
                    streak.current_streak = 1
                streak.longest_streak = max(streak.longest_streak, streak.current_streak)
                streak.last_played_date = today_str

        db.commit()

        return {
            "success": True,
            "streak": {
                "current": streak.current_streak,
                "longest": streak.longest_streak,
                "total_games": streak.total_games,
            },
        }
    finally:
        db.close()


@app.get("/api/games/leaderboard")
async def games_leaderboard():
    today = date.today()
    today_str = today.isoformat()
    monday = today - timedelta(days=today.weekday())
    monday_str = monday.isoformat()
    db = SessionLocal()
    try:
        week_results = db.query(GameResult).filter(GameResult.game_date >= monday_str).all()
        today_results = [r for r in week_results if r.game_date == today_str]

        # Aggregate weekly scores
        weekly: dict = {}
        for r in week_results:
            weekly.setdefault(r.recruiter_email, {"weekly_score": 0})
            weekly[r.recruiter_email]["weekly_score"] += r.score

        # Aggregate today's stats
        today_agg: dict = {}
        for r in today_results:
            today_agg.setdefault(r.recruiter_email, {"times": [], "count": 0})
            today_agg[r.recruiter_email]["count"] += 1
            if r.time_taken and r.time_taken > 0:
                today_agg[r.recruiter_email]["times"].append(r.time_taken)

        leaderboard = []
        for email, wdata in weekly.items():
            rec = db.query(Recruiter).filter(Recruiter.email == email).first()
            streak_rec = db.query(Streak).filter(Streak.recruiter_email == email).first()
            td = today_agg.get(email, {"times": [], "count": 0})
            times = td["times"]
            avg_time = round(sum(times) / len(times)) if times else 0
            fastest = min(times) if times else 0
            leaderboard.append({
                "email": email,
                "name": rec.name if rec else email,
                "weekly_score": wdata["weekly_score"],
                "games_played_today": td["count"],
                "current_streak": streak_rec.current_streak if streak_rec else 0,
                "avg_time_today": avg_time,
                "fastest_game_today": fastest,
            })

        leaderboard.sort(key=lambda x: (-x["weekly_score"], x["avg_time_today"] or 9999))
        return leaderboard[:10]
    finally:
        db.close()


@app.get("/api/games/my-stats")
async def my_game_stats(email: str):
    db = SessionLocal()
    try:
        streak = db.query(Streak).filter(Streak.recruiter_email == email).first()
        total_score = sum(r.score for r in db.query(GameResult).filter(GameResult.recruiter_email == email).all())
        return {
            "current_streak": streak.current_streak if streak else 0,
            "longest_streak": streak.longest_streak if streak else 0,
            "total_games": streak.total_games if streak else 0,
            "total_score": total_score,
        }
    finally:
        db.close()


# ─── Health ───────────────────────────────────────────────────────────────────

@app.get("/health")
async def health():
    return {"status": "ok"}
